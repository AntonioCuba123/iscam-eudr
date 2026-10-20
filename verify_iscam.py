# -*- coding: utf-8 -*-
import openpyxl, os
wb = openpyxl.load_workbook(r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx', data_only=False)
ws_i = wb['ISCAM_Score']
ws_m = wb['MATRIZ_174']
ws_sc = wb['SCORECARDS']

print('=== ESTRUCTURA DE FILAS ===')
for r in range(1, 6):
    v = ws_i.cell(r, 1).value
    print(f'  ISCAM_Score row{r}: {str(v)[:50] if v else "(vacio)"}')

print()
print('=== HEADERS ROW 3 (fila de labels) ===')
for col in [1, 5, 6, 7, 8, 16, 21, 26, 27, 28, 29]:
    v = ws_i.cell(3, col).value
    if v: print(f'  col{col}: {str(v)[:40]}')

print()
print('=== PRIMER DATUM (row 4) ===')
print(f'  RUC ISCAM row4: {ws_i.cell(4,1).value}')
print(f'  RUC MATRIZ row3: {ws_m.cell(3,1).value}')
print(f'  Coinciden: {str(ws_i.cell(4,1).value) == str(ws_m.cell(3,1).value)}')

print()
print('=== FORMULAS CLAVE ===')
print(f'  D2 F4: {ws_i.cell(4,6).value}')
print(f'  D3 G4: {ws_i.cell(4,7).value}')
print(f'  IP1 P4: {ws_i.cell(4,16).value}')
print(f'  IP Score Z4: {ws_i.cell(4,26).value}')
print(f'  IV Score AA4: {ws_i.cell(4,27).value}')
print(f'  RE AB4: {ws_i.cell(4,28).value}')
print(f'  Tipo AC4: {str(ws_i.cell(4,29).value)[:60]}')

print()
print('=== SCORECARDS REFS ===')
print(f'  SC!C3 (IP): {ws_sc.cell(3,3).value}')
print(f'  SC!N3 (Tipo): {ws_sc.cell(3,14).value}')

data_rows = [r for r in range(4, ws_i.max_row+1)
             if ws_i.cell(r,1).value and len(str(ws_i.cell(r,1).value))==11]
print()
print(f'Data rows ISCAM: {len(data_rows)} (rows 4-{data_rows[-1]})')
print(f'Ultimo RUC ISCAM: {ws_i.cell(data_rows[-1],1).value}')
print(f'Ultimo RUC MATRIZ: {ws_m.cell(176,1).value}')

sz = os.path.getsize(r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx')
print()
print(f'Archivo: {sz:,} bytes')
print('Hojas:', wb.sheetnames)

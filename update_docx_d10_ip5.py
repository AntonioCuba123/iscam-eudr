# -*- coding: utf-8 -*-
"""Actualiza el docx: D10 = bandera pura e IP5 con la certificación por conteo 0/33/66/100."""
import os, shutil, copy, datetime
import docx
from docx.oxml.ns import qn

ROOT = r"C:\Users\acuba\eudrperu"
PATH = os.path.join(ROOT, "ISCAM", "Metodologia - Instrumento 2 e Indice ISCAM.docx")
ts = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
bak = os.path.join(ROOT, "ISCAM", "respaldos", f"Metodologia - Instrumento 2 e Indice ISCAM.backup_{ts}_pre-d10-bandera.docx")
shutil.copy2(PATH, bak); print("respaldo:", bak)

d = docx.Document(PATH)

def edit_run(p, old, new, tag):
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)
            print(f"OK {tag}")
            return True
    raise SystemExit(f"NO ENCONTRADO en {tag}: {old[:80]!r}")

# ── 1) §2.6 — qué promedia ──────────────────────────────────────────────────
edit_run(d.paragraphs[68],
         "(D1\u2013D10; D11 es inventario y no promedia, y de D10 solo punt\u00faa la asesor\u00eda EUDR, "
         "porque su certificaci\u00f3n y su oferta de \u00abcertificaci\u00f3n EUDR\u00bb son banderas)",
         "(D1\u2013D9; D10 es bandera y D11 inventario: ninguna de las dos promedia y sus \u00edtems se "
         "reportan aparte)",
         "[068] \u00a72.6 qu\u00e9 promedia")

# ── 2) Tabla 1 · fila IP5 ───────────────────────────────────────────────────
edit_run(d.tables[1].rows[6].cells[1].paragraphs[0],
         "= (D7 + D9 + Certificaci\u00f3n) / 3, con la certificaci\u00f3n llevada a 0\u2013100 "
         "(el score 0\u20131 se multiplica por 100) para que las tres componentes compartan escala",
         "= (D7 + D9 + Certificaci\u00f3n) / 3, donde Certificaci\u00f3n es el conteo de certificaciones "
         "voluntarias vigentes llevado a 0\u201333\u201366\u2013100 (0 \u00b7 1 \u00b7 2 \u00b7 \u22653 esquemas), "
         "para que las tres componentes compartan la escala 0\u2013100",
         "T1 r6 IP5")

# ── 3) §5 hito — IP5 corregido ──────────────────────────────────────────────
edit_run(d.paragraphs[89],
         "IP5 qued\u00f3 adem\u00e1s corregido de escala: sumaba D7 y D9 (0\u2013100) con la certificaci\u00f3n "
         "en 0\u20131; la certificaci\u00f3n entra ahora multiplicada por 100, de modo que las tres "
         "componentes est\u00e1n en la misma escala 0\u2013100.",
         "IP5 qued\u00f3 adem\u00e1s corregido de escala y de fuente: sumaba D7 y D9 (0\u2013100) con la "
         "certificaci\u00f3n en 0\u20131, y tomaba la certificaci\u00f3n del score ponderado precensal "
         "(0,6 org\u00e1nica / 0,4 Comercio Justo). Ahora usa el conteo documentado 0\u201333\u201366\u2013100 "
         "de certificaciones vigentes \u2014B0.9 al cargar la encuesta, registro precensal mientras "
         "tanto\u2014, de modo que las tres componentes comparten la escala 0\u2013100; y devuelve "
         "\u00ab\u2014\u00bb si falta D7, D9 o la certificaci\u00f3n, nunca 0.",
         "[089] \u00a75 IP5")

# ── 4) §5 hito nuevo — D10 bandera pura (se inserta tras el hito ISA-PV) ────
src = d.paragraphs[86]._p                       # "ISA-PactoVerde provisional (2026-10-15)"
new_p = copy.deepcopy(src)
runs = new_p.findall(qn('w:r'))
for extra in runs[1:]:
    new_p.remove(extra)
t = runs[0].find(qn('w:t'))
t.text = ("D10 ratificado como bandera pura (2026-10-15). Sus tres \u00edtems \u2014falsa seguridad sobre "
          "la certificaci\u00f3n (D10.36), oferta de \u00abcertificaci\u00f3n EUDR\u00bb (D10.37) y asesor\u00eda "
          "recibida (D10.36a)\u2014 se capturan, se codifican y se reportan aparte; ninguno promedia. El "
          "promedio de D10 queda vac\u00edo y la dimensi\u00f3n no entra a ning\u00fan \u00edndice (D10 no alimenta "
          "ISA-EUDR, IP ni IV). Fuente: \u00a77.3 de la metodolog\u00eda ISA \u2014\u00abD10 \u00b7 Certificaci\u00f3n "
          "(bandera) \u2026 Bandera: no promedia\u00bb\u2014 y Tabla 4 del documento de trabajo.")
t.set(qn('xml:space'), 'preserve')
d.paragraphs[87]._p.addprevious(new_p)
print("OK [nuevo] \u00a75 hito D10 bandera pura")

d.save(PATH)
print("docx guardado:", PATH)

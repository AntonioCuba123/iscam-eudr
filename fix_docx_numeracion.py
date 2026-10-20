# -*- coding: utf-8 -*-
"""Reintegra a la numeración dos viñetas que habían quedado sueltas al agregarlas
con python-docx (párrafo 42, §2.2 —ratificación— y párrafo 85, §5 —hito de la
ratificación—) y les fija la tipografía del resto del documento (Arial 11)."""
import os, copy, datetime, shutil
import docx
from docx.shared import Pt
from docx.oxml.ns import qn

ROOT = r"C:\Users\acuba\eudrperu"
DOC = os.path.join(ROOT, "ISCAM", "Metodologia - Instrumento 2 e Indice ISCAM.docx")
ts = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")
shutil.copy2(DOC, os.path.join(ROOT, "ISCAM", "respaldos",
             f"Metodologia - Instrumento 2 e Indice ISCAM.backup_{ts}_pre-numeracion.docx"))

d = docx.Document(DOC)


def numpr_of(p):
    ppr = p._p.find(qn('w:pPr'))
    return ppr.find(qn('w:numPr')) if ppr is not None else None


def adopt(target_idx, model_idx):
    t, m = d.paragraphs[target_idx], d.paragraphs[model_idx]
    mn = numpr_of(m)
    if mn is None:
        print('  (modelo sin numPr)', model_idx); return
    tppr = t._p.find(qn('w:pPr'))
    if tppr is None:
        tppr = t._p.makeelement(qn('w:pPr'), {})
        t._p.insert(0, tppr)
    old = tppr.find(qn('w:numPr'))
    if old is not None:
        tppr.remove(old)
    tppr.append(copy.deepcopy(mn))
    for r in t.runs:
        if r.font.name is None:
            r.font.name = 'Arial'
        if r.font.size is None:
            r.font.size = Pt(11)
    print(f'  p{target_idx} ← numPr de p{model_idx}',
          [r.font.name for r in t.runs], '|', t.text[:60])


print('p41 numPr:', numpr_of(d.paragraphs[41]) is not None,
      '| p43 numPr:', numpr_of(d.paragraphs[43]) is not None)
adopt(42, 41 if numpr_of(d.paragraphs[41]) is not None else 43)
adopt(85, 86)

d.save(DOC)
print('OK — numeración y tipografía uniformadas.')

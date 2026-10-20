# -*- coding: utf-8 -*-
"""Construye la MATRIZ FINAL: datos internos (BASE CONSOLIDADA) + hallazgos web por RUC."""
import pandas as pd

FINDINGS = {}
DIR = r"C:\Users\acuba\eudrperu\BASE CAFE CACAO 174"
for f in ["findings_parte1.py", "findings_parte2.py", "findings_parte3.py", "findings_parte4.py",
          "findings_parte5.py", "findings_parte6.py", "findings_parte7.py", "findings_parte8.py",
          "findings_parte9.py", "findings_parte10.py"]:
    with open(DIR + "\\" + f, encoding="utf-8") as fh:
        exec(fh.read(), globals())

print("RUC con hallazgos web:", len(FINDINGS))

dim_org = pd.read_excel(r"C:\Users\acuba\eudrperu\ficharucsunat\BASE CONSOLIDADA.xlsx",
                        sheet_name="DIM_ORGANIZACION")
dim_org["ruc"] = dim_org["ruc"].astype(str).str.zfill(11)

rucs_df = pd.read_excel(r"C:\Users\acuba\eudrperu\BASE CAFE CACAO 174\CONSULTA MASIVA RUC 10.10.26\ruces_lotes_174.xlsx")
rucs_df["RUC"] = rucs_df["RUC"].astype(str).str.zfill(11)

m = rucs_df.merge(dim_org, left_on="RUC", right_on="ruc", how="left")

cols_int = ["RUC", "razon_social", "nombre_comercial", "departamento", "provincia", "distrito",
            "producto_dominante", "estado_contribuyente", "condicion_contribuyente",
            "fecha_inicio_actividades", "inscrito_rnp", "exporta", "fob_total_usd",
            "n_socios_varones", "n_socias_mujeres", "hectareas_totales",
            "certificacion_organica", "certificacion_comercio_justo", "avance_eudr",
            "buen_contribuyente", "agente_retencion", "presidente_vigente",
            "gerente_general", "n_duas", "score_completitud"]
m = m[[c for c in cols_int if c in m.columns]]

m["socios_total_interno"] = pd.to_numeric(m.get("n_socios_varones"), errors="coerce").fillna(0) + \
                            pd.to_numeric(m.get("n_socias_mujeres"), errors="coerce").fillna(0)
m["socios_total_interno"] = m["socios_total_interno"].replace(0, pd.NA)

def g(ruc, k):
    d = FINDINGS.get(ruc) or {}
    v = d.get(k)
    return "" if v in (None, "No verificado") else v

m["nombre_publico_web"] = m["RUC"].map(lambda r: g(r, "nombre_publico"))
m["web_oficial"] = m["RUC"].map(lambda r: g(r, "web"))
m["socios_web"] = m["RUC"].map(lambda r: g(r, "socios"))
m["hectareas_web"] = m["RUC"].map(lambda r: g(r, "hectareas"))
m["certificaciones_web"] = m["RUC"].map(lambda r: g(r, "certificaciones"))
m["exportaciones_web"] = m["RUC"].map(lambda r: g(r, "exportaciones"))
m["hallazgo_web"] = m["RUC"].map(lambda r: g(r, "hallazgo_clave"))
m["fuentes_web"] = m["RUC"].map(lambda r: g(r, "fuentes"))
m["tiene_hallazgo_web"] = m["RUC"].map(lambda r: "SI" if r in FINDINGS else "NO")
m["origen_hallazgo"] = m["RUC"].map(
    lambda r: "WEB PÚBLICA (buscador, con fuentes)" if r in FINDINGS else "PENDIENTE DE BÚSQUEDA")

final_cols = ["RUC", "razon_social", "nombre_publico_web", "nombre_comercial", "departamento",
              "provincia", "distrito", "producto_dominante", "estado_contribuyente",
              "condicion_contribuyente", "fecha_inicio_actividades", "socios_total_interno",
              "socios_web", "hectareas_totales", "hectareas_web", "certificacion_organica",
              "certificacion_comercio_justo", "certificaciones_web", "exporta", "fob_total_usd",
              "exportaciones_web", "avance_eudr", "buen_contribuyente", "agente_retencion",
              "inscrito_rnp", "presidente_vigente", "gerente_general", "n_duas",
              "tiene_hallazgo_web", "origen_hallazgo", "hallazgo_web", "fuentes_web",
              "web_oficial", "score_completitud"]
final = m[[c for c in final_cols if c in m.columns]]

final = final.rename(columns={
    "RUC": "RUC", "razon_social": "Razón Social (SUNAT)",
    "nombre_publico_web": "Nombre público / comercial (web)",
    "nombre_comercial": "Nombre comercial (interno)",
    "departamento": "Departamento", "provincia": "Provincia", "distrito": "Distrito",
    "producto_dominante": "Producto dominante", "estado_contribuyente": "Estado SUNAT",
    "condicion_contribuyente": "Condición SUNAT",
    "fecha_inicio_actividades": "Inicio actividades",
    "socios_total_interno": "Socios (interno)", "socios_web": "Socios (web)",
    "hectareas_totales": "Hectáreas (interno)", "hectareas_web": "Hectáreas (web)",
    "certificacion_organica": "Cert. Orgánica (interno)",
    "certificacion_comercio_justo": "Cert. Comercio Justo (interno)",
    "certificaciones_web": "Certificaciones (web)", "exporta": "Exporta (interno)",
    "fob_total_usd": "FOB total USD (interno)", "exportaciones_web": "Exportaciones (web)",
    "avance_eudr": "Avance EUDR", "buen_contribuyente": "Buen Contribuyente",
    "agente_retencion": "Agente Retención", "inscrito_rnp": "Inscrito RNP",
    "presidente_vigente": "Presidente vigente", "gerente_general": "Gerente general",
    "n_duas": "N° DUAs", "tiene_hallazgo_web": "¿Tiene hallazgo web?",
    "origen_hallazgo": "Origen del hallazgo",
    "hallazgo_web": "HALLAZGO DETALLADO (fuentes públicas)",
    "fuentes_web": "FUENTES CITADAS (URLs)", "web_oficial": "Web / referencia principal",
    "score_completitud": "Score completitud (interno)"})

final = final.sort_values("RUC").reset_index(drop=True)

out_csv = r"C:\Users\acuba\eudrperu\BASE CAFE CACAO 174\MATRIZ_ORGANIZACIONES_174_FINAL.csv"
out_xlsx = r"C:\Users\acuba\eudrperu\BASE CAFE CACAO 174\CONSULTA MASIVA RUC 10.10.26\MATRIZ_ORGANIZACIONES_174_FINAL.xlsx"

final.to_csv(out_csv, index=False, encoding="utf-8-sig")
with pd.ExcelWriter(out_xlsx, engine="openpyxl") as w:
    final.to_excel(w, sheet_name="MATRIZ 174", index=False)
    ws = w.sheets["MATRIZ 174"]
    from openpyxl.styles import Font, Alignment, PatternFill
    hdr_fill = PatternFill("solid", fgColor="1F4E78")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(vertical="center", wrap_text=True)
    widths = {"A": 13, "B": 38, "C": 34, "D": 20, "E": 14, "F": 14, "G": 14, "H": 12, "I": 12,
              "J": 13, "K": 13, "L": 11, "M": 26, "N": 12, "O": 22, "P": 12, "Q": 14, "R": 30,
              "S": 11, "T": 14, "U": 34, "V": 12, "W": 13, "X": 13, "Y": 11, "Z": 20, "AA": 20,
              "AB": 9, "AC": 14, "AD": 26, "AE": 90, "AF": 70, "AG": 40, "AH": 12}
    for i, wd in widths.items():
        ws.column_dimensions[i].width = wd
    ws.freeze_panes = "C2"
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.auto_filter.ref = ws.dimensions

print("Filas:", len(final))
print("Con hallazgo web:", (final["¿Tiene hallazgo web?"] == "SI").sum())
print("Pendientes:", (final["¿Tiene hallazgo web?"] == "NO").sum())
print("Columnas:", len(final.columns))
print("CSV :", out_csv)
print("XLSX:", out_xlsx)
pend = final[final["¿Tiene hallazgo web?"] == "NO"]["RUC"].tolist()
print("RUC pendientes:", pend)

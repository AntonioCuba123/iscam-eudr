# -*- coding: utf-8 -*-
"""Actualiza header de D12 en MATRIZ_174 con nota metodológica clara."""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

SRC = r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE.xlsx'

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb['MATRIZ_174']

# D12 está en col W = 23
COL = 23

# Header fila 1 (nombre corto)
c1 = ws.cell(1, COL)
c1.value = 'D12'
c1.fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
c1.font = Font(bold=True)
c1.alignment = Alignment(horizontal='center')

# Header fila 2 (descripción completa con nota)
c2 = ws.cell(2, COL)
c2.value = (
    'D12 — Huella Carbono\n'
    '0=Sin medición\n'
    '1=En curso (ha conocidas)\n'
    '2=Certificado\n'
    '─────────────────────\n'
    'NOTA: Valor actual = proxy FAO/IPCC\n'
    '(ha × factor IPCC/FAO por cultivo).\n'
    'Reemplazar tras Instrumento 2 Q40-Q44.'
)
c2.fill = PatternFill(start_color='BDD7EE', end_color='BDD7EE', fill_type='solid')
c2.font = Font(bold=True, size=8)
c2.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)

# Ajustar altura de fila 2 para que se vea completa
ws.row_dimensions[2].height = 100

wb.save(SRC)
print(f'Header D12 actualizado en {SRC}')
print('  Fila 1: D12')
print('  Fila 2: descripción + nota FAO/IPCC proxy')

# -*- coding: utf-8 -*-
"""Documenta la re-ponderación de ISA-PV en el Word (§2.4) y en PONDERADORES."""
import copy, os
import docx
from docx.text.paragraph import Paragraph
import win32com.client as win32

ROOT = r"C:\Users\acuba\eudrperu"
DOC = os.path.join(ROOT, "ISCAM", "Metodologia - Instrumento 2 e Indice ISCAM.docx")
XLS = os.path.join(ROOT, "PLANTILLA_INDICE_PVE_ISCAM.xlsx")

TXT = ("ISA-PactoVerde aplica la misma regla que ISA-EUDR: se pondera solo sobre los componentes con "
       "dato y la suma de ponderaciones se re-normaliza a 100%. Un componente sin dato ya no cuenta como "
       "cero. PPWR (sacos reutilizables, peso 4%) queda excluido mientras no tenga fuente de captura, y su "
       "cobertura se reporta aparte (MATRIZ_174, columna AQ).")

# ── Word ──────────────────────────────────────────────────────────────────
doc = docx.Document(DOC)
ancla = None
for p in doc.paragraphs:
    if p.text.strip().startswith("Regla operativa en la hoja de cálculo"):
        ancla = p
        break
if ancla is None:
    raise SystemExit("no encontré el ancla en §2.4")

if "ISA-PactoVerde aplica la misma regla" in "\n".join(x.text for x in doc.paragraphs):
    print("Word: la nota ya estaba, no duplico")
else:
    nuevo = copy.deepcopy(ancla._p)
    ancla._p.addnext(nuevo)
    np = Paragraph(nuevo, ancla._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    np.runs[0].text = TXT
    doc.save(DOC)
    print("Word: nota de re-ponderación ISA-PV insertada en §2.4")

# ── Excel: nota en PONDERADORES ───────────────────────────────────────────
xl = win32.Dispatch("Excel.Application")
xl.Visible = False
xl.DisplayAlerts = False
xl.ScreenUpdating = False
xl.EnableEvents = False
wb = xl.Workbooks.Open(XLS)
xl.Calculation = -4135
pond = wb.Worksheets("PONDERADORES")
# f29 = ISA-PactoVerde
pond.Cells(29, 2).Value = (str(pond.Cells(29, 2).Value) +
                           "  Desde 2026-10-15 se REPONDERA sobre los componentes presentes "
                           "(un componente sin dato ya no cuenta como 0); PPWR queda excluido mientras no "
                           "tenga fuente. Ver MATRIZ_174!AP y la cobertura en MATRIZ_174!AQ.")
last = 56
while pond.Cells(last, 1).Value is not None and last < 70:
    last += 1
pond.Cells(last, 1).Value = "ISA-PV REPONDERADA (2026-10-15)"
pond.Cells(last, 1).Font.Bold = True
pond.Cells(last + 1, 1).Value = ("ISA-PV = Σ(wᵢ·vᵢ) ÷ Σwᵢ sobre los componentes con dato, como ISA-EUDR. "
                                 "Antes el denominador fijo (100) hacía que un componente sin dato contara "
                                 "como 0 y recortaba el índice.")
pond.Cells(last + 1, 1).WrapText = True
xl.Calculation = -4105
xl.CalculateFull()
wb.Save()
wb.Close(SaveChanges=True)
xl.Quit()
print(f"Excel: notas escritas en PONDERADORES B29 y A{last}:A{last+1}")
print("PDF pendiente de regenerar")

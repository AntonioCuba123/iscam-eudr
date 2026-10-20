# -*- coding: utf-8 -*-
import docx
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
for i in (55, 60, 68, 87, 88, 89, 90):
    print(f'--- [{i:03d}] ---')
    print(d.paragraphs[i].text)
print()
for ti in (1, 4, 5):
    tb = d.tables[ti]
    print(f'===== TABLA {ti} ({len(tb.rows)}x{len(tb.columns)}) =====')
    for ri, row in enumerate(tb.rows):
        print(f'  r{ri}:', ' || '.join(c.text.replace('\n', ' ') for c in row.cells))

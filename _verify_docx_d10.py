# -*- coding: utf-8 -*-
import docx
from docx.oxml.ns import qn
from collections import Counter
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
print('parrafos', len(d.paragraphs), 'tablas', len(d.tables))
for i in range(83, 93):
    p = d.paragraphs[i]
    npr = p._p.find(qn('w:pPr'))
    num = None
    if npr is not None:
        n = npr.find(qn('w:numPr'))
        if n is not None:
            num = 'num'
    print(f'[{i:03d}] style={p.style.name!r} {num} :: {p.text[:160]}')
print()
print('--- busqueda de referencias viejas ---')
full = '\n'.join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            full += '\n' + c.text
for old in ('solo puntúa la asesoría', 'solo asesoria', 'multiplica por 100)',
            'D1–D10; D11 es inventario', 'score 0–1 se multiplica'):
    print(f'  {old!r}: {"PRESENTE" if old in full else "ausente"}')
print()
print('--- fuentes usadas ---')
fonts = Counter()
for p in d.paragraphs:
    for r in p.runs:
        if r.text.strip():
            fonts[(r.font.name, r.font.size)] += 1
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    if r.text.strip():
                        fonts[(r.font.name, r.font.size)] += 1
print(fonts)

# -*- coding: utf-8 -*-
import docx
from docx.oxml.ns import qn
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
for i, p in enumerate(d.paragraphs):
    if 'D10' in p.text or 'D11' in p.text:
        print(f'[{i:03d}]', p.text[:500]); print()
print('=== parrafos con runs sin fuente explicita ===')
for i, p in enumerate(d.paragraphs):
    if any(r.font.name is None and r.text.strip() for r in p.runs):
        print(f'[{i:03d}]', repr(p.text[:110]))
print()
print('=== [068] completo ===')
print(d.paragraphs[68].text)
print()
print('=== [087] nuevo, completo ===')
print(d.paragraphs[87].text)
print()
print('=== [090] completo ===')
print(d.paragraphs[90].text)

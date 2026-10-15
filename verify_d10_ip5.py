# -*- coding: utf-8 -*-
import openpyxl, warnings, zipfile, re
from collections import Counter
warnings.filterwarnings('ignore')
p = r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx'
wbv = openpyxl.load_workbook(p, data_only=True)
wbf = openpyxl.load_workbook(p)

print("hojas:", wbv.sheetnames)

n = wbv['NORMALIZADO']
print("NORMALIZADO!CK (D10) distintos:", Counter(str(n.cell(r, 89).value) for r in range(3, 177)))

m = wbv['MATRIZ_174']
au = Counter(str(m.cell(r, 47).value) for r in range(3, 177))
av = Counter(str(m.cell(r, 48).value) for r in range(3, 177))
print("MATRIZ_174!AU (n cert) distintos:", au)
print("MATRIZ_174!AV (score) distintos:", av)
print("AU2/AV2:", repr(m.cell(2, 47).value), '|', repr(m.cell(2, 48).value))

s = wbv['ISCAM_Score']
ts = Counter(str(s.cell(r, 20).value) for r in range(4, 178))
print("ISCAM_Score!T (IP5) distintos:", ts)
print("N3:", repr(s.cell(3, 14).value))
print("A192/A193:", repr(s.cell(192, 1).value), '|', repr(s.cell(193, 1).value))
print("N4 formula:", wbf['ISCAM_Score'].cell(4, 14).value)
print("T4 formula:", wbf['ISCAM_Score'].cell(4, 20).value)
print("T5 formula:", wbf['ISCAM_Score'].cell(5, 20).value)

c = wbv['SCORECARDS']
print("SCORECARDS!L distintos:", Counter(str(c.cell(r, 12).value) for r in range(3, 177)))

r = wbv['RADAR_BRECHAS']
print("RADAR K2/L2:", repr(r.cell(2, 11).value), repr(r.cell(2, 12).value), "| K3/L3:", repr(r.cell(3, 11).value), repr(r.cell(3, 12).value))

# errores de fórmula
errs = []
for ws in wbv.worksheets:
    for row in ws.iter_rows():
        for cc in row:
            if isinstance(cc.value, str) and re.match(r'^#(REF|VALUE|DIV|NAME|N/A|NUM|NULL)', cc.value):
                errs.append((ws.title, cc.coordinate, cc.value))
print("errores de formula:", len(errs), errs[:10])

# validaciones listas (raw xml)
z = zipfile.ZipFile(p)
xml = z.read('xl/worksheets/sheet4.xml').decode('utf-8', 'ignore')  # CAPTURA
print("CAPTURA dataValidation list:", xml.count('type="list"'), "| whole:", xml.count('type="whole"'))
print("AU1/AV1 vacías:", m.cell(1, 47).value, m.cell(1, 48).value)

# -*- coding: utf-8 -*-
"""Agrega al Excel la hoja RESUMEN VERIFICACION con los datos reales hallados en vivo."""
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

XLSX = (r"C:\Users\acuba\eudrperu\BASE CAFE CACAO 174"
        r"\CONSULTA MASIVA RUC 10.10.26\ENLACES_BUSQUEDA_DUA_2024-2026.xlsx")
CSV = r"C:\Users\acuba\eudrperu\ficharucsunat\_verif_dua_links.csv"

v = pd.read_csv(CSV, dtype=str)
v["conteo"] = pd.to_numeric(v["conteo"], errors="coerce").fillna(0).astype(int)
v["ruc"] = v["ident"].str.extract(r"RUC=(\d+)")
v["anio"] = v["ident"].str.extract(r"Año=(\d+)")

razon = dict(zip(pd.read_excel(XLSX, sheet_name="1. Resumen por año", dtype=str)["RUC"],
                 pd.read_excel(XLSX, sheet_name="1. Resumen por año", dtype=str)["Razón social"]))

n1 = v[v["nivel"] == "n1"]
n2 = v[v["nivel"] == "n2"]
ok1 = n1[n1["estado"] == "OK"]
ok2 = n2[n2["estado"] == "OK"]

anios_por_ruc = ok1.groupby("ruc")["anio"].nunique().to_dict()
combos_por_ruc = ok1.groupby("ruc")["conteo"].sum().to_dict()
duas_por_ruc = ok2.groupby("ruc")["conteo"].sum().to_dict()
links1_ok = ok1.groupby("ruc").size().to_dict()
links2_ok = ok2.groupby("ruc").size().to_dict()

rucs = sorted(set(v["ruc"].dropna()))
filas = []
for ruc in rucs:
    a = anios_por_ruc.get(ruc, 0)
    c = combos_por_ruc.get(ruc, 0)
    d = duas_por_ruc.get(ruc, 0)
    if d > 0:
        sit = "EXPORTA - DUAs confirmadas"
    elif c > 0:
        sit = "EXPORTA - combos sin DUA (revisar)"
    elif a > 0:
        sit = "EXPORTA - resumen con datos"
    else:
        sit = "SIN MOVIMIENTO 2024-2026"
    filas.append({
        "RUC": ruc, "Razón social": razon.get(ruc, ""),
        "Situación detectada": sit,
        "Años con datos (de 3)": a,
        "Años con datos": ", ".join(sorted(ok1[ok1["ruc"] == ruc]["anio"].unique())) or "-",
        "Combos en resúmenes": c,
        "DUAs confirmadas (nivel 2)": d,
        "Enlaces nivel 1 OK": links1_ok.get(ruc, 0),
        "Enlaces nivel 2 OK": links2_ok.get(ruc, 0),
    })
df = pd.DataFrame(filas)

orden = {"EXPORTA - DUAs confirmadas": 0, "EXPORTA - combos sin DUA (revisar)": 1,
         "EXPORTA - resumen con datos": 2, "SIN MOVIMIENTO 2024-2026": 3}
df = df.sort_values(["Situación detectada", "DUAs confirmadas (nivel 2)"],
                    key=lambda s: s.map(orden) if s.name == "Situación detectada" else s,
                    ascending=[True, False]).reset_index(drop=True)

conteo = df["Situación detectada"].value_counts().to_dict()
guia = [("RESUMEN DE LA VERIFICACIÓN EN VIVO", ""),
        ("", ""),
        ("Fuente", "ADUANET SUNAT - aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias "
                   "(régimen 40 = exportación definitiva)"),
        ("Fecha", "11 de octubre de 2026"),
        ("Enlaces verificados", "%d (nivel 1: %d · nivel 2: %d)"
         % (len(v), len(n1), len(n2))),
        ("Responden con datos", "%d (%.0f%%)"
         % (int((v['estado'] == 'OK').sum()), 100 * (v['estado'] == 'OK').mean())),
        ("Vacíos", "%d" % int((v['estado'] == 'VACIO').sum())),
        ("Errores de red", "%d" % int((v['estado'] == 'ERROR').sum())),
        ("Tiempo de verificación", "123 segundos con 4 workers concurrentes"),
        ("", ""),
        ("HALLAZGOS POR ORGANIZACIÓN", ""),
        ("RUC con DUAs confirmadas", str(conteo.get("EXPORTA - DUAs confirmadas", 0))),
        ("RUC con combos pero sin DUA", str(conteo.get("EXPORTA - combos sin DUA (revisar)", 0))),
        ("RUC con resumen con datos", str(conteo.get("EXPORTA - resumen con datos", 0))),
        ("RUC sin movimiento 2024-2026", str(conteo.get("SIN MOVIMIENTO 2024-2026", 0))),
        ("Total de DUAs confirmadas", str(int(ok2["conteo"].sum()))),
        ("Total de combos en resúmenes", str(int(ok1["conteo"].sum()))),
        ("", ""),
        ("LECTURA", ""),
        ("Exportador registrado ≠ exportador activo",
         "%d de los 76 RUC figuran como EXPORTADOR o IMPORTADOR/EXPORTADOR en la ficha SUNAT, "
         "pero %d no registran movimiento en régimen 40 durante 2024-2026. Puede deberse a "
         "exportaciones bajo otro régimen, venta a través de terceros ( comercializadora ), o "
         "inactividad." % (len(rucs), conteo.get("SIN MOVIMIENTO 2024-2026", 0))),
        ("Canal", "Todas las DUAs listadas salen en canal VERDE en las consultas realizadas."),
        ("Nivel 2 muy confiable", "%d de %d enlaces de listado devolvieron DUAs: los candidatos "
         "armados desde el FOB previo son un atajo válido."
         % (len(ok2), len(n2))),
        ("", ""),
        ("HOJAS DEL LIBRO", ""),
        ("1. Resumen por año", "228 enlaces con estado de verificación y N° de combos"),
        ("2. Listado DUAs", "471 enlaces con estado, N° de DUAs y detalle (aduana/n°/fecha/canal)"),
        ("3. DUAs ya capturadas", "1141 enlaces de detalle ya resueltos en el proyecto"),
        ("4. Estado por RUC", "Ruta de búsqueda por organización"),
        ("5. Resumen verificación", "Esta hoja: situación real de cada RUC"),
        ]
dfg = pd.DataFrame(guia, columns=["Campo", "Detalle"])

wb = load_workbook(XLSX)
for h in ("RESUMEN VERIFICACION",):
    if h in wb.sheetnames:
        del wb[h]
ws = wb.create_sheet("RESUMEN VERIFICACION")

hf = PatternFill("solid", fgColor="1F4E78")
verde = PatternFill("solid", fgColor="C6EFCE")
ambar = PatternFill("solid", fgColor="FFEB9C")
gris = PatternFill("solid", fgColor="F2F2F2")
rojo = PatternFill("solid", fgColor="FFC7CE")
thin = Side(style="thin", color="BFBFBF")

# bloque guía
for i, (_, r) in enumerate(dfg.iterrows(), start=1):
    a = ws.cell(row=i, column=1, value=r["Campo"])
    b = ws.cell(row=i, column=2, value=r["Detalle"])
    if i > 1 and r["Campo"] and str(r["Campo"]).isupper() and not r["Detalle"]:
        a.font = Font(bold=True, color="1F4E78", size=11)
    b.alignment = Alignment(vertical="top", wrap_text=True)
    a.alignment = Alignment(vertical="top", wrap_text=True)

inicio = len(dfg) + 3
COL = {"EXPORTA - DUAs confirmadas": (verde, "1F6F3D", True),
       "EXPORTA - combos sin DUA (revisar)": (ambar, "9C6500", False),
       "EXPORTA - resumen con datos": (ambar, "9C6500", False),
       "SIN MOVIMIENTO 2024-2026": (gris, "808080", False)}

for j, c in enumerate(df.columns, start=1):
    cell = ws.cell(row=inicio, column=j, value=c)
    cell.font = Font(bold=True, color="FFFFFF", size=10)
    cell.fill = hf
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
ws.row_dimensions[inicio].height = 34

for i, (_, r) in enumerate(df.iterrows(), start=inicio + 1):
    for j, c in enumerate(df.columns, start=1):
        cell = ws.cell(row=i, column=j, value=r[c])
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        if c == "Situación detectada":
            bg, fg, bold = COL[r[c]]
            cell.fill = bg
            cell.font = Font(bold=bold, size=9, color=fg)

anchos = [14, 46, 30, 10, 22, 12, 14, 12, 12]
for i, a in enumerate(anchos[:len(df.columns)], start=1):
    ws.column_dimensions[get_column_letter(i)].width = a
ws.column_dimensions["A"].width = 40
ws.column_dimensions["B"].width = 46
ws.freeze_panes = ws.cell(row=inicio + 1, column=3)

# mover la hoja al final
wb.move_sheet("RESUMEN VERIFICACION", offset=len(wb.sheetnames) - 1 - wb.sheetnames.index("RESUMEN VERIFICACION"))
wb.save(XLSX)

print("Hoja RESUMEN VERIFICACION agregada")
print("Hojas:", wb.sheetnames)
print()
print(df["Situación detectada"].value_counts().to_string())
print()
print("Top 12 RUC por DUAs confirmadas:")
print(df.head(12)[["RUC", "Razón social", "Situación detectada",
                   "Años con datos (de 3)", "DUAs confirmadas (nivel 2)"]].to_string(index=False))

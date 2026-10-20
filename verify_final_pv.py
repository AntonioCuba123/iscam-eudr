# -*- coding: utf-8 -*-
"""Verificación final (no destructiva) del libro entregable."""
import openpyxl, collections, os, zipfile, re

PATH = r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx'
wb = openpyxl.load_workbook(PATH, data_only=True)
wf = openpyxl.load_workbook(PATH, data_only=False)

print('hojas:', wb.sheetnames)
print('tamaño:', f'{os.path.getsize(PATH):,}', 'bytes')
print()

# errores de fórmula en valores cacheados
errs = collections.Counter()
for sn in wb.sheetnames:
    for row in wb[sn].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('#'):
                errs[c.value] += 1
print('errores de fórmula:', dict(errs) or 'NINGUNO')

# validaciones (x14) preservadas
z = zipfile.ZipFile(PATH)
x = z.read('xl/worksheets/sheet4.xml').decode('utf-8', errors='ignore')
print('validaciones CAPTURA: list =', len(re.findall(r'<x14:dataValidation[^>]*>.*?</x14:dataValidation>', x, re.S)),
      '| la referencia contiene type="list"',
      len([m for m in re.findall(r'<x14:dataValidation[^>]*>', x) if 'list' in m]),
      '| whole =', len(re.findall(r'<dataValidation type="whole"', x)))

cap = wb['CAPTURA']
print('CAPTURA:', cap.max_row, 'x', cap.max_column)

# --- ISA-PV ---
m = wb['MATRIZ_174']
est = collections.Counter()
snum = 0
for r in range(3, 177):
    est[m.cell(r, 46).value] += 1
    if isinstance(m.cell(r, 19).value, (int, float)):
        snum += 1
print()
print('MATRIZ_174!AT (estado ISA-PV):', dict(est))
print('MATRIZ_174!S con valor numérico:', snum, 'de 174')
print('SCORECARDS!B186 (PVE promedio):', wb['SCORECARDS']['B186'].value)
print('ISCAM_Score!AM (estado) muestra:', [wb['ISCAM_Score'].cell(r, 39).value for r in (4, 5, 6)])

# --- IP ---
si = wb['ISCAM_Score']
print()
print('ISCAM_Score IP1..IP5 (P4:T4) valores:', [si.cell(4, c).value for c in range(16, 21)])
print('ISCAM_Score IP4 (Z4):', si.cell(4, 26).value, '| IV (AA4):', si.cell(4, 27).value,
      '| RE (AB4):', si.cell(4, 28).value)
print('ISCAM_Score!T4 formula:', wf['ISCAM_Score']['T4'].value)
print('ISCAM_Score!Q4 formula:', wf['ISCAM_Score']['Q4'].value)
print('ISCAM_Score!AA4 formula:', wf['ISCAM_Score']['AA4'].value)

# --- W:Z ---
print()
print('MATRIZ_174!W3 formula:', wf['MATRIZ_174']['W3'].value)
print('MATRIZ_174!X3 formula:', wf['MATRIZ_174']['X3'].value)
cnt = collections.Counter((m.cell(r, 23).value, m.cell(r, 24).value, m.cell(r, 25).value, m.cell(r, 26).value)
                          for r in range(3, 177))
print('MATRIZ W/X/Y/Z valores:', dict(cnt))

# --- cabeceras / notas ---
print()
print('MATRIZ AP2 :', wf['MATRIZ_174']['AP2'].value[:80], '...')
print('MATRIZ AT2 :', repr(wf['MATRIZ_174']['AT2'].value))
print('MATRIZ T2  :', repr(wf['MATRIZ_174']['T2'].value)[:90])
print('MATRIZ A182:', wf['MATRIZ_174']['A182'].value)
print('ISCAM AI3  :', repr(wf['ISCAM_Score']['AI3'].value))
print('ISCAM AM3  :', repr(wf['ISCAM_Score']['AM3'].value))
print('POND A58   :', wf['PONDERADORES']['A58'].value)
print('INSTR A5   :', wf['INSTRUCCIONES']['A5'].value[:90], '...')
print('SCORE A186 :', wf['SCORECARDS']['A186'].value)
print('MATRIZ T3/T4/T5:', [m.cell(r, 20).value for r in (3, 4, 5)])
print('MATRIZ AP3/AP4/AP5:', [m.cell(r, 42).value for r in (3, 4, 5)])

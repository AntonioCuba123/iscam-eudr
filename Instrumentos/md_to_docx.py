"""
Convierte DOC_TRABAJO_Instrumento2_ISCAM_v2.md a .docx
usando python-docx con estilos académicos.
"""

import re
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import copy

# ── helpers ────────────────────────────────────────────────────────────────

def set_cell_background(cell, hex_color):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)

def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    tblBorders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{side}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), 'AAAAAA')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def apply_inline_markup(paragraph, text):
    """Aplica negritas, cursivas y código inline dentro de un párrafo."""
    # Combinar negritas y cursivas: ***texto***
    pattern = re.compile(r'(\*\*\*(.+?)\*\*\*|\*\*(.+?)\*\*|_(.+?)_|\*(.+?)\*|`(.+?)`|🚩\s*\*\*(.+?)\*\*)')
    last = 0
    for m in pattern.finditer(text):
        # texto antes del match
        if m.start() > last:
            paragraph.add_run(text[last:m.start()])
        s = m.group(0)
        if s.startswith('***'):
            r = paragraph.add_run(m.group(2))
            r.bold = True
            r.italic = True
        elif s.startswith('**'):
            # detectar bandera 🚩 **...**
            inner = m.group(3)
            r = paragraph.add_run(inner)
            r.bold = True
        elif s.startswith('_') or s.startswith('*'):
            inner = m.group(4) or m.group(5)
            r = paragraph.add_run(inner)
            r.italic = True
        elif s.startswith('`'):
            r = paragraph.add_run(m.group(6))
            r.font.name = 'Courier New'
            r.font.size = Pt(9)
        elif '🚩' in s:
            r0 = paragraph.add_run('🚩 ')
            inner = m.group(7)
            r = paragraph.add_run(inner)
            r.bold = True
        last = m.end()
    if last < len(text):
        paragraph.add_run(text[last:])

def add_paragraph(doc, text, style='Normal', alignment=None, space_before=0, space_after=6):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    if alignment:
        p.alignment = alignment
    apply_inline_markup(p, text)
    return p

def add_code_block(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Cm(1)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    # fondo gris claro via shading en el párrafo
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'F0F0F0')
    pPr.append(shd)
    return p

def parse_table_md(lines):
    """Extrae filas de una tabla markdown (lista de líneas)."""
    rows = []
    for line in lines:
        line = line.strip()
        if re.match(r'^\|[-:| ]+\|$', line):
            continue  # separador
        if line.startswith('|') and line.endswith('|'):
            cells = [c.strip() for c in line[1:-1].split('|')]
            rows.append(cells)
    return rows

def add_table(doc, rows):
    if not rows:
        return
    ncols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = 'Table Grid'
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        for c_idx, cell_text in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.paragraphs[0].paragraph_format.space_before = Pt(1)
            cell.paragraphs[0].paragraph_format.space_after = Pt(1)
            if r_idx == 0:
                # encabezado
                set_cell_background(cell, 'E8EDF5')
                run_holder = cell.paragraphs[0]
                apply_inline_markup(run_holder, cell_text)
                for run in run_holder.runs:
                    run.bold = True
                    run.font.size = Pt(8.5)
            else:
                set_cell_background(cell, 'FFFFFF' if r_idx % 2 == 0 else 'F9FAFB')
                p = cell.paragraphs[0]
                apply_inline_markup(p, cell_text)
                for run in p.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph()  # espacio después de tabla

# ── estilos del documento ──────────────────────────────────────────────────

def setup_document():
    doc = Document()

    # Márgenes
    for section in doc.sections:
        section.top_margin    = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin   = Cm(3.0)
        section.right_margin  = Cm(2.5)

    # Fuente base
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing = Pt(14)
    style.paragraph_format.space_after  = Pt(6)

    # H1
    h1 = doc.styles['Heading 1']
    h1.font.name  = 'Calibri'
    h1.font.size  = Pt(16)
    h1.font.bold  = True
    h1.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6E)
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after  = Pt(8)

    # H2
    h2 = doc.styles['Heading 2']
    h2.font.name  = 'Calibri'
    h2.font.size  = Pt(13)
    h2.font.bold  = True
    h2.font.color.rgb = RGBColor(0x1F, 0x3B, 0x6E)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after  = Pt(6)

    # H3
    h3 = doc.styles['Heading 3']
    h3.font.name   = 'Calibri'
    h3.font.size   = Pt(11)
    h3.font.bold   = True
    h3.font.italic = False
    h3.font.color.rgb = RGBColor(0x2E, 0x5C, 0x99)
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after  = Pt(4)

    # List Bullet
    try:
        lb = doc.styles['List Bullet']
        lb.font.name = 'Calibri'
        lb.font.size = Pt(11)
    except Exception:
        pass

    return doc

# ── conversor principal ────────────────────────────────────────────────────

def convert(md_path, docx_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    doc = setup_document()

    i = 0
    in_table   = False
    table_rows = []
    in_code    = False
    code_lines = []
    in_note    = False  # bloques > Note

    def flush_table():
        nonlocal in_table, table_rows
        if table_rows:
            add_table(doc, table_rows)
        in_table   = False
        table_rows = []

    def flush_code():
        nonlocal in_code, code_lines
        if code_lines:
            add_code_block(doc, '\n'.join(code_lines))
        in_code    = False
        code_lines = []

    while i < len(lines):
        raw  = lines[i]
        line = raw.rstrip('\n')
        stripped = line.strip()

        # ── bloque de código ──────────────────────────────────────────────
        if stripped.startswith('```'):
            if not in_code:
                in_code    = True
                code_lines = []
            else:
                flush_code()
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        # ── tabla ─────────────────────────────────────────────────────────
        if stripped.startswith('|'):
            in_table = True
            table_rows.append([c.strip() for c in stripped.strip('|').split('|')])
            # saltar separador en la siguiente línea si es ---
            if i + 1 < len(lines) and re.match(r'^\|[-:| ]+\|', lines[i+1].strip()):
                i += 2
                continue
            i += 1
            continue
        else:
            if in_table:
                flush_table()

        # ── línea en blanco ───────────────────────────────────────────────
        if stripped == '':
            in_note = False
            i += 1
            continue

        # ── separador horizontal ──────────────────────────────────────────
        if re.match(r'^---+$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after  = Pt(6)
            run = p.add_run()
            # línea horizontal via borde inferior
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '6')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), 'AAAAAA')
            pBdr.append(bottom)
            pPr.append(pBdr)
            i += 1
            continue

        # ── encabezados ───────────────────────────────────────────────────
        hm = re.match(r'^(#{1,4})\s+(.*)', stripped)
        if hm:
            level = len(hm.group(1))
            text  = hm.group(2)
            style_map = {1: 'Heading 1', 2: 'Heading 2', 3: 'Heading 3', 4: 'Heading 3'}
            p = doc.add_paragraph(style=style_map.get(level, 'Heading 3'))
            apply_inline_markup(p, text)
            i += 1
            continue

        # ── lista con viñeta ──────────────────────────────────────────────
        bullet = re.match(r'^[-*]\s+(.*)', stripped)
        if bullet:
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.left_indent  = Cm(0.8)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(2)
            apply_inline_markup(p, bullet.group(1))
            i += 1
            continue

        # ── lista numerada ────────────────────────────────────────────────
        numbered = re.match(r'^\d+\.\s+(.*)', stripped)
        if numbered:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.left_indent  = Cm(0.8)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after  = Pt(2)
            apply_inline_markup(p, numbered.group(1))
            i += 1
            continue

        # ── bloque de nota / cita (>) ─────────────────────────────────────
        if stripped.startswith('>'):
            note_text = re.sub(r'^>\s*', '', stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent  = Cm(1.0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after  = Pt(2)
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            left = OxmlElement('w:left')
            left.set(qn('w:val'), 'single')
            left.set(qn('w:sz'), '8')
            left.set(qn('w:space'), '8')
            left.set(qn('w:color'), '6B90C2')
            pBdr.append(left)
            pPr.append(pBdr)
            apply_inline_markup(p, note_text)
            for run in p.runs:
                run.font.size  = Pt(9.5)
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # ── texto normal ──────────────────────────────────────────────────
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(6)
        apply_inline_markup(p, stripped)
        i += 1

    # flush final
    if in_table:
        flush_table()
    if in_code:
        flush_code()

    doc.save(docx_path)
    print(f'Guardado: {docx_path}')

# ── main ───────────────────────────────────────────────────────────────────

if __name__ == '__main__':
    import os
    base = r'c:\Users\acuba\eudrperu\Instrumentos'
    convert(
        os.path.join(base, 'DOC_TRABAJO_Instrumento2_ISCAM_v2.md'),
        os.path.join(base, 'DOC_TRABAJO_Instrumento2_ISCAM_v2.docx')
    )

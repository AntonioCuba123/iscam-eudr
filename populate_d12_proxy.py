# -*- coding: utf-8 -*-
"""Enriquece D12 con proxy compuesto: certificaciones + FOB UE."""

import openpyxl
from openpyxl.styles import PatternFill, Alignment

SRC  = r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE.xlsx'
COL  = 23  # D12

GRIS     = PatternFill(start_color='D9D9D9', end_color='D9D9D9', fill_type='solid')
AMARILLO = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
VERDE_L  = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
CENTER   = Alignment(horizontal='center')

wb = openpyxl.load_workbook(SRC, data_only=False)
ws = wb['MATRIZ_174']

cambios = []
for r in range(3, ws.max_row + 1):
    cell = ws.cell(r, COL)
    ha   = ws.cell(r, 18).value or 0
    cert = str(ws.cell(r, 8).value or '').lower().strip()
    fair = str(ws.cell(r, 9).value or '').lower().strip()
    fob_ue = ws.cell(r, 12).value or 0
    org  = ws.cell(r, 2).value or ''

    # ¿Tiene certificación?
    tiene_cert = (
        any(x in cert for x in ['si', 'sí', 'ue', 'organic', 'cer', 'flo', 'fai', 'fa', 'vig', 'en '])
        or any(x in fair for x in ['si', 'sí', 'fair', 'flo', 'fai', 'fa', 'vig', 'en'])
    )

    old_val = cell.value

    if ha > 0:
        # Ya estaba como 1 (proxy FAO) — se mantiene
        nuevo = 1
        razon = f'ha={ha} (proxy FAO)'
    elif tiene_cert:
        # Tiene certificacion sin ha — alto indicio de medicion
        nuevo = 1
        razon = 'certificacion detectada'
    elif fob_ue > 500000:
        # Alto FOB UE sin cert ni ha — sophisticacion + presion UE
        nuevo = 1
        razon = f'FOB UE={fob_ue:>12.0f} (proxy presion UE)'
    else:
        nuevo = 0
        razon = 'sin datos — pendiente Inst.2'

    cell.value = nuevo
    if nuevo == 1:
        cell.fill = AMARILLO
    else:
        cell.fill = GRIS
    cell.alignment = CENTER

    if old_val != nuevo:
        cambios.append(f'  {str(org)[:40]:40} | {old_val}->{nuevo} | {razon}')

print(f'=== Cambios en D12 ({len(cambios)} filas modificadas) ===')
for c in cambios:
    print(c)

wb.save(SRC)
print(f'\nD12 enriquecido guardado en {SRC}')

# Resumen
wb2 = openpyxl.load_workbook(SRC, data_only=True)
ws2 = wb2['MATRIZ_174']
scores = {}
for r in range(3, ws2.max_row+1):
    d = ws2.cell(r, COL).value
    scores[d] = scores.get(d, 0) + 1
print('\nResumen D12:')
for k in sorted(scores.keys(), reverse=True):
    print(f'  Score {k}: {scores[k]} orgs')

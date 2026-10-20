# -*- coding: utf-8 -*-
import openpyxl, warnings, re, zipfile
warnings.filterwarnings('ignore')
p = r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx'
wb = openpyxl.load_workbook(p)

s = wb['SCORECARDS']
print('SCORECARDS fila 2 cols A..P:')
from openpyxl.utils import get_column_letter as gl
for c in range(1, 17):
    print('  ', gl(c), repr(str(s.cell(2, c).value)[:60]), '| r3:', str(s.cell(3, c).value)[:70])
print()
print('referencias a SCORECARDS!L / \$L:')
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and re.search(r'SCORECARDS!\$?L\d', v):
                print('  ', ws.title, c.coordinate, '|', v[:150])
print()
cap = wb['CAPTURA']
print('CAPTURA!N2 =', repr(cap.cell(2, 14).value))
print('CAPTURA!N3..N6 =', [cap.cell(r, 14).value for r in range(3, 7)])
print('CAPTURA columnas CV/CW/CX (D10):', [(gl(x), cap.cell(2, x).value) for x in (100, 101, 102)])
# raw validations on sheet CAPTURA
z = zipfile.ZipFile(p)
names = [n for n in z.namelist() if 'worksheets/sheet' in n]
print('sheets xml:', names)

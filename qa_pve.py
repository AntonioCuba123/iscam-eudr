# -*- coding: utf-8 -*-
"""QA de las fórmulas de PLANTILLA_INDICE_PVE.xlsx"""
import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE.xlsx")
print("=== HOJAS ===")
for name in wb.sheetnames:
    ws = wb[name]
    print(f"  {name}: {ws.max_row} filas x {ws.max_column} cols")

print("\n=== HOJA: PONDERADORES ===")
ws = wb["PONDERADORES"]
# Fila 12-18: componentes
for r in range(12, 19):
    v = ws.cell(row=r, column=4).value
    print(f"  Fila {r}: D{r} = {v}")
# Total
total = ws.cell(row=20, column=4).value
print(f"  TOTAL (fila 20): {total}")

print("\n=== HOJA: MATRIZ_174 ===")
ws = wb["MATRIZ_174"]
# Fila 3: primera org
print(f"  RUC fila 3: {ws.cell(3,1).value}")
print(f"  FOB fila 3: {ws.cell(3,5).value}")
print(f"  % UE (fórmula col M=13): {ws.cell(3,13).value}")
print(f"  Score Cert (fórmula col O=15): {ws.cell(3,15).value}")

print("\n=== HOJA: PREGUNTAS_PROXY ===")
ws = wb["PREGUNTAS_PROXY"]
# Buscar una pregunta con fórmula de score
for r in range(3, 80):
    v = ws.cell(row=r, column=5).value
    if v and isinstance(v, str) and v.startswith("="):
        print(f"  Fila {r} score: {v[:80]}...")
        break

print("\n=== HOJA: SCORECARDS ===")
ws = wb["SCORECARDS"]
print(f"  RUC fila 3: {ws.cell(3,1).value}")
pve = ws.cell(3, 3).value
print(f"  PVE fórmula (fila 3): {pve}")
sem = ws.cell(3, 4).value
print(f"  Semáforo (fila 3): {sem}")
rank = ws.cell(3, 5).value
print(f"  Ranking (fila 3): {rank}")
seg = ws.cell(3, 14).value
print(f"  Segmento (fila 3): {seg}")
# Resumen
for r in range(ws.max_row - 15, ws.max_row + 1):
    v1 = ws.cell(row=r, column=1).value
    v2 = ws.cell(row=r, column=2).value
    if v1:
        print(f"  Resumen fila {r}: {v1} = {v2}")

print("\n=== HOJA: RADAR_BRECHAS ===")
ws = wb["RADAR_BRECHAS"]
print(f"  Fila 3 (SAN MARTIN) D2 fórmula: {ws.cell(3,3).value}")
print(f"  PVE promedio SAN MARTIN: {ws.cell(3,13).value}")

print("\n=== VALIDACIONES CLAVE ===")
errors = []
# 1. Total ponderadores debe sumar 100
ws_p = wb["PONDERADORES"]
total = ws_p.cell(row=20, column=4).value
if isinstance(total, str) and total.startswith("=SUM"):
    errors.append(f"  ⚠ Total ponderadores es fórmula (pendiente recalc): {total}")
else:
    errors.append(f"  ✅ Total ponderadores: {total}")

# 2. PVE fórmula en SCORECARDS debe ser correcta
ws_s = wb["SCORECARDS"]
pve = ws_s.cell(3, 3).value
if pve and pve.startswith("=") and "PONDERADORES" in pve:
    errors.append(f"  ✅ PVE fórmula incluye PONDERADORES refs")
else:
    errors.append(f"  ⚠ PVE fórmula: {pve}")

# 3. Semáforo en SCORECARDS
sem = ws_s.cell(3, 4).value
if sem and ("Crítico" in sem or "Critico" in sem):
    errors.append(f"  ✅ Semáforo con etiquetas correctas")
else:
    errors.append(f"  ⚠ Semáforo: {sem}")

for e in errors:
    print(e)

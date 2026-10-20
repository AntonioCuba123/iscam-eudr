# -*- coding: utf-8 -*-
"""
Normaliza el idioma de la documentación del proyecto a castellano.

Regla aplicada: se traduce el anglicismo cuando aparece en prosa española y
existe equivalente castellano de uso normal (score → puntaje, performance →
desempeño, ranking → ordenamiento, dashboard → panel, gap → brecha, etc.).

Se CONSERVAN intactos (no son errores de idioma):
  · nombres propios e identificadores: ISCAM_Score, SCORECARDS, RADAR_BRECHAS,
    nombres de archivo (`Checklist - …`), scripts (`update_docx_*.py`), variables;
  · citas literales y títulos de publicaciones (van entre comillas/itálica);
  · URLs, código y palabras clave SQL (JOIN, HEADER);
  · el abstract en inglés de los working papers (convención académica);
  · términos naturalizados en el español académico/técnico: paper, proxy,
    software, know-how, toolkit (nombre del producto), Test Exportador
    (herramienta de PromPerú), Global GAP (certificación), Fairtrade, etc.;
  · siglas y nombres de instrumentos: EUDR, DDS, FSA, CSDDD, LMR, CalCafé…

Además corrige defectos de idioma heredados del portapapeles: caracteres CJK,
cirílicos y kana intrusos en texto español (文献, 满分, 報告, список, 第一部,
法令, 理论上, 跨国, 允许, に立ち, 生成) y erratas tipográficas (poblad,
referencing, severely).

Idempotente: una segunda corrida no produce cambios.
"""
import os, re, shutil, datetime, glob
import docx

ROOT = os.path.dirname(os.path.abspath(__file__))
RESP = os.path.join(ROOT, "ISCAM", "respaldos")
ts = datetime.datetime.now().strftime("%Y-%m-%d_%H%M")

# ── 1) reglas ──────────────────────────────────────────────────────────
# 1a. Frases exactas (se aplican antes que las reglas de palabra).
FRASES = [
    # — reparación de concordancias de artículo (reemplazos genéricos previos) —
    ('la desempeño', 'el desempeño'),
    ('La desempeño', 'El desempeño'),
    ('el brecha', 'la brecha'),
    ('El brecha', 'La brecha'),
    ('un brecha', 'una brecha'),
    ('del fecha límite', 'de la fecha límite'),
    ('al fecha límite', 'a la fecha límite'),
    # — título de publicación en inglés: se restaura y se protege —
    ('EUDR preparación assessment', 'EUDR readiness assessment'),
    # — defectos de scripts ajenas y erratas —
    (' estructura de跨国 corporation', ' estructura de corporación transnacional'),
    ('Se，允许 el uso', 'Se permite el uso'),
    ('presentに立ち responsAbility', 'presentó responsAbility'),
    ('Ricardo |生成:', 'Ricardo | Generado:'),
    (' список para consulta OSCE masiva', ' lista para consulta OSCE masiva'),
    ('derechos第一部', 'derechos'),
    ('del 法令 de Comunidades Nativas', 'de la Ley de Comunidades Nativas'),
    (':**理论上 son', ':** en teoría son'),
    ('文献', 'bibliografía'),
    ('con满分 en trazabilidad', 'con un puntaje máximo en trazabilidad'),
    ('報告ar', 'reportar'),
    ('severely restringida', 'severamente restringida'),
    ('poblad en', 'poblados en'),
    ('referencing SCORECARDS', 'referenciados desde SCORECARDS'),
    # — checklist / roadmap / health check —
    ('Checklist de ', 'Lista de verificación de '),
    ('Checklist — ', 'Lista de verificación — '),
    ('la Checklist (', 'la lista de verificación ('),
    ('(checklist `', '(lista de verificación `'),
    ('del checklist', 'de la lista de verificación'),
    ('checklist', 'lista de verificación'),
    ('Roadmap DT-ISCAM-001', 'Hoja de ruta DT-ISCAM-001'),
    ('health check', 'verificación de salud'),
    # — score: compuestos que suenan mejor con el nombre primero —
    ('PVE Score', 'Puntaje PVE'),
    ('D3 Score', 'Puntaje D3'),
    ('Gender score', 'puntaje de género'),
    ('gender score', 'puntaje de género'),
    ('Score trazabilidad', 'Puntaje de trazabilidad'),
    ('score trazabilidad', 'puntaje de trazabilidad'),
    ('Score de certificaciones', 'Puntaje de certificaciones'),
    ('Score principal', 'Puntaje principal'),
    ('Score promedio', 'Puntaje promedio'),
    ('Score Cert', 'Puntaje Cert'),
    ('Scorecards por organización', 'Tarjetas de puntaje por organización'),
    # — default / input / join / gender —
    ('el default recomendado', 'el valor predeterminado recomendado'),
    ('del default de este documento', 'del valor predeterminado de este documento'),
    ('un default,', 'un valor predeterminado,'),
    ('Input hardcoded', 'Entrada fija'),
    ('Input editable', 'Entrada editable'),
    ('es el input que define', 'es la entrada que define'),
    ('clave de join', 'clave de unión'),
    ('join por RUC', 'unión por RUC'),
    # — gap / flag / baseline —
    ('Índice de Readiness Cooperativo Dashboard por organización',
     'Panel por organización del índice de preparación cooperativa'),
    ('El flag ', 'El indicador '),
    ('RUC, flag registrado', 'RUC, indicador de registro'),
    ('el gap real', 'la brecha real'),
    ('El gap real', 'La brecha real'),
    ('gap de ISA-EUDR', 'brecha de ISA-EUDR'),
    ('el gap de diseño', 'la brecha de diseño'),
    ('gap de «% volumen trazable»', 'brecha de «% volumen trazable»'),
    ('Los Gaps Reales', 'Las brechas reales'),
    ('Feb sin baseline', 'Feb sin línea base'),
    ('del framework', 'del marco'),
    ('email', 'correo'),
]

# 1b. Reglas de palabra completa (con preservación de mayúscula inicial).
GENERICAS = {
    'scores': 'puntajes', 'score': 'puntaje',
    'performance': 'desempeño',
    'ranking': 'ordenamiento',
    'dashboard': 'panel',
    'online': 'en línea',
    'gender': 'género',
    'deadline': 'fecha límite',
    'readiness': 'preparación',
    'baseline': 'línea base',
    'gaps': 'brechas', 'gap': 'brecha',
    'flag': 'indicador',
}
WORD_RE = re.compile(r'\b(' + '|'.join(GENERICAS) + r')\b', re.I)

# 1c. Tramos que nunca se tocan.
PROTEGIDO = [
    r'https?://\S+',                 # URLs
    r'`[^`\n]*`',                    # código/identificadores en línea
    r'\*"[^"\n]*"\*',                # citas literales en itálica
    r'\bGlobal GAP\b', r'\bGAP\b',   # certificación
    r'\bTest Exportador\b',          # herramienta PromPerú
    r'\bGender & inclusion\b',       # criterio citado del estudio SAI
    r'the gap is enormous',          # cita textual en inglés
    r'\bJOIN\b', r'\bHEADER\b',      # SQL
    r'EUDR readiness assessment[^\n]*',   # título de publicación (EN)
]
PROT_RE = re.compile('|'.join(PROTEGIDO))

_CONTADOR = {'frases': 0, 'palabras': 0}


def _aplica(seg):
    if not seg:
        return seg
    for a, b in FRASES:
        if a in seg:
            _CONTADOR['frases'] += seg.count(a)
            seg = seg.replace(a, b)

    def rep(m):
        w = m.group(1)
        esp = GENERICAS[w.lower()]
        _CONTADOR['palabras'] += 1
        return esp.capitalize() if w[:1].isupper() else esp

    return WORD_RE.sub(rep, seg)


def normaliza(txt):
    """Devuelve el texto con los anglicismos de prosa traducidos."""
    if not txt:
        return txt
    out, pos = [], 0
    for m in PROT_RE.finditer(txt):
        out.append(_aplica(txt[pos:m.start()]))
        out.append(m.group(0))
        pos = m.end()
    out.append(_aplica(txt[pos:]))
    return ''.join(out)


# ── 2) markdown ────────────────────────────────────────────────────────
DIRS_MD = ['.', 'Instrumentos', 'ISCAM', 'reports', 'docs', 'ficharucsunat',
           'EUDR_ISCAM_Referencias', 'Informes']
EXCLUIR = ('respaldos', '.git', '__pycache__', 'eda_output', 'Scripts',
           'appsperu', 'cache_ocds', 'EUDR\\', 'EUDR/', 'adex cein',
           'BASE CAFE CACAO', 'fichaduanas')


def listar(ext, dirs):
    out = []
    for d in dirs:
        base = ROOT if d == '.' else os.path.join(ROOT, d)
        for p in glob.glob(os.path.join(base, '**', '*' + ext), recursive=True):
            rel = os.path.relpath(p, ROOT)
            if any(x in rel for x in EXCLUIR):
                continue
            out.append(p)
    return sorted(set(out))


def proc_md(path):
    with open(path, encoding='utf-8') as f:
        original = f.read()
    nuevo = normaliza(original)
    if nuevo != original:
        with open(path, 'w', encoding='utf-8') as f:
            f.write(nuevo)
        return True
    return False


# ── 3) docx ────────────────────────────────────────────────────────────
def _fix_parrafo(p, en_abstract):
    if en_abstract or not p.runs:
        return 0
    full = ''.join(r.text for r in p.runs)
    nuevo = normaliza(full)
    if nuevo == full:
        return 0
    n = 0
    for r in p.runs:                       # intento por run (conserva formato)
        t2 = normaliza(r.text)
        if t2 != r.text:
            r.text = t2
            n += 1
    if ''.join(r.text for r in p.runs) == nuevo:
        return n
    p.runs[0].text = nuevo                 # consolidación si el texto cruza runs
    for r in p.runs[1:]:
        r.text = ''
    return 1


def proc_docx(path):
    bak = os.path.join(RESP, os.path.basename(path) + f'.backup_{ts}_pre-idioma.docx')
    shutil.copy2(path, bak)
    d = docx.Document(path)
    cambios = 0
    en_abstract = False
    for p in d.paragraphs:
        t = p.text.strip()
        if t == 'Abstract':
            en_abstract = True
        elif t.lower().startswith('palabras clave'):
            en_abstract = False
        cambios += _fix_parrafo(p, en_abstract)
    for tb in d.tables:
        for row in tb.rows:
            for cell in row.cells:
                en_abstract = False
                for p in cell.paragraphs:
                    cambios += _fix_parrafo(p, en_abstract)
    if cambios:
        d.save(path)
    return cambios, bak


# ── 4) main ────────────────────────────────────────────────────────────
if __name__ == '__main__':
    print('=== MARKDOWN ===')
    mds = listar('.md', DIRS_MD)
    cambiados_md = []
    for p in mds:
        if proc_md(p):
            cambiados_md.append(os.path.relpath(p, ROOT))
    print(f'  revisados: {len(mds)} | modificados: {len(cambiados_md)}')
    for f in cambiados_md:
        print('   ', f)

    print('=== DOCX ===')
    docxs = listar('.docx', DIRS_MD)
    for p in docxs:
        cambios, bak = proc_docx(p)
        if cambios:
            print(f'  {os.path.relpath(p, ROOT)}: {cambios} párrafos/celdas')

    print('=== XLSX (se procesan con fix_idioma_xlsx.py) ===')
    print(f"  frases aplicadas: {_CONTADOR['frases']} | palabras: {_CONTADOR['palabras']}")

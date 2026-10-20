# -*- coding: utf-8 -*-
import docx
from docx.oxml.ns import qn
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
for i in range(78, 96):
    p = d.paragraphs[i]
    numpr = p._p.find(qn('w:pPr'))
    num = None
    if numpr is not None:
        npr = numpr.find(qn('w:numPr'))
        if npr is not None:
            ilvl = npr.find(qn('w:ilvl')); nid = npr.find(qn('w:numId'))
            num = (ilvl.get(qn('w:val')) if ilvl is not None else None,
                   nid.get(qn('w:val')) if nid is not None else None)
    print(f'[{i:03d}] style={p.style.name!r} numPr={num} text={p.text[:70]!r}')
print()
print('--- parrafos 40-44 (ratificacion) ---')
for i in range(40, 45):
    p = d.paragraphs[i]
    print(f'[{i:03d}] style={p.style.name!r} text={p.text[:90]!r}')

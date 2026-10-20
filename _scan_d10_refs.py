# -*- coding: utf-8 -*-
import openpyxl, warnings, re
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook(r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx')

pats = [r'NORMALIZADO!CK\b', r'NORMALIZADO!CL\b', r'ISCAM_Score!N\b', r'ISCAM_Score!\$N',
        r'\$CK\d', r'ISCAM_Score!N\d', r'!N\d']
rx = re.compile('|'.join(pats))
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and rx.search(v):
                print(ws.title, c.coordinate, '|', v[:200])

print('=== bloque de cobertura ISCAM_Score (A170:C205) ===')
s = wb['ISCAM_Score']
for r in range(170, 206):
    vals = [str(s.cell(r, cc).value)[:70] for cc in (1, 2, 3)]
    if any(v != 'None' for v in vals):
        print(r, vals)

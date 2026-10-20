# -*- coding: utf-8 -*-
import docx
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
for i in (68, 89):
    print(f'--- [{i}] runs ---')
    for j, r in enumerate(d.paragraphs[i].runs):
        print(f'  run{j}: bold={r.bold} font={r.font.name} size={r.font.size} :: {r.text[:120]!r}')
print('--- T1 r6 ---')
for j, r in enumerate(d.tables[1].rows[6].cells[1].paragraphs[0].runs):
    print(f'  run{j}: bold={r.bold} font={r.font.name} size={r.font.size} :: {r.text[:160]!r}')
print('--- T5 r5 c1 ---')
for j, r in enumerate(d.tables[5].rows[5].cells[1].paragraphs[0].runs):
    print(f'  run{j}: bold={r.bold} font={r.font.name} size={r.font.size} :: {r.text[:160]!r}')
print('--- T4 r4 c2 ---')
for j, r in enumerate(d.tables[4].rows[4].cells[2].paragraphs[0].runs):
    print(f'  run{j}: bold={r.bold} font={r.font.name} size={r.font.size} :: {r.text[:160]!r}')

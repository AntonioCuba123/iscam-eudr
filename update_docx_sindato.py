# -*- coding: utf-8 -*-
"""Alinea el documento con el libro: umbral ISA-PV, regla «sin dato» precensal y coberturas."""
import copy
import os

import docx
from docx.text.paragraph import Paragraph

DOC = r"C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx"
doc = docx.Document(DOC)


def reemplaza(prefijo, nuevo):
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefijo):
            runs = p.runs
            if not runs:
                return False
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            runs[0].text = nuevo
            return True
    return False


def inserta_despues(prefijo, texto):
    ancla = None
    for p in doc.paragraphs:
        if p.text.strip().startswith(prefijo):
            ancla = p
            break
    if ancla is None:
        return False
    nuevo = copy.deepcopy(ancla._p)
    ancla._p.addnext(nuevo)
    np = Paragraph(nuevo, ancla._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    np.runs[0].text = texto
    return True


# 1) §2.4 umbral ISA-PV: ahora implementado en el libro
ok1 = reemplaza(
    "Umbral mínimo ISA-PactoVerde",
    "Umbral mínimo ISA-PactoVerde: 3 de las 4 dimensiones Pacto Verde (D12–D15) con dato; por debajo del "
    "umbral, ISA-PactoVerde y el promedio PVE se reportan como «—» (sin dato), nunca como 0 —un valor "
    "calculado sobre una sola dimensión sería una medición falsa—. Por encima del umbral el índice se "
    "computa solo sobre los componentes con dato, re-normalizando las ponderaciones. El umbral se evalúa "
    "sobre la captura (hoja NORMALIZADO), no sobre el proxy precensal.")

# 2) §2.4 nueva viñeta: regla «sin dato» sobre la base precensal
TXT = ("Regla sobre la base precensal (MATRIZ_174, 2026-10-15): los ceros que provenían de la falta de dato "
       "ya no se muestran como 0. La certificación se recodificó a SI/NO tomando el texto fuente —venía "
       "truncada a tres caracteres por el generador de la matriz—, y el score de certificación, el % UE/FOB, "
       "el FOB total y el de la Unión Europea y los indicadores P1–P14 aparecen como «—» cuando falta la "
       "fuente. Un cero medido —por ejemplo, una organización que exporta pero no a la Unión Europea— se "
       "conserva. Los índices derivados (IV1, IV4 e IP5) quedan también sin dato cuando su insumo lo está.")
ok2 = inserta_despues("ISA-PactoVerde aplica la misma regla que ISA-EUDR", TXT)

# 3) §5 coberturas corregidas
ok3 = reemplaza(
    "Capa de ingesta (implementada 2026-10-15)",
    "Capa de ingesta (implementada 2026-10-15). La tabla de codificación, la hoja de captura y la "
    "normalización están construidas y cableadas al motor (hojas CODIFICACION, CAPTURA y NORMALIZADO). Lo "
    "que resta es trabajo de campo: cargar las respuestas en CAPTURA. Cobertura de partida: D3 (legalidad "
    "OSCE/OEFA) 174/174; D1, D2 y D4–D11 sin dato hasta aplicar el instrumento; comercio exterior (% UE/FOB "
    "y FOB UE) 52/174; certificación verificada 22/174; D12–D15 por captura 0/174 —el proxy FAO/IPCC vigente "
    "se conserva en MATRIZ_174 y se reemplaza al levantar el Instrumento 2—. La consulta de legalidad "
    "OSCE/OEFA está completa para las 174 organizaciones (2026-10-15): 0 inhabilitaciones vigentes y 1 "
    "infracción ambiental OEFA.")

# 4) §5 IV1–IV5: sin dato cuando falta el insumo
ok4 = reemplaza(
    "Reescribir IV1–IV5",
    "Reescribir IV1–IV5 con las definiciones de diseño una vez ingresado el instrumento (insumos ya "
    "capturados); IV4 aplica certificación (70/30) + agravante de legalidad 30% × (100 − D3). Mientras la "
    "certificación o el comercio exterior no tengan fuente, IV4 e IV1 se reportan como «—», no como 0: los "
    "60 y 70 del diseño son supuestos de valor, no umbrales de ausencia.")

for n, ok in [("umbral ISA-PV", ok1), ("viñeta precensal", ok2), ("coberturas", ok3), ("IV1-IV5", ok4)]:
    print(("OK   " if ok else "FALLA"), n)

doc.save(DOC)
print("documento guardado")

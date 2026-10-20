# -*- coding: utf-8 -*-
"""Actualiza el documento de metodología con la capa de ingesta implementada.

Uso:  restaura el .docx desde el respaldo _pre-ingesta y vuelve a aplicar todos
los cambios (script idempotente).
"""
import copy, glob, os, shutil, sys
import docx
from docx.text.paragraph import Paragraph

ROOT = r"C:\Users\acuba\eudrperu"
DOC = os.path.join(ROOT, "ISCAM", "Metodologia - Instrumento 2 e Indice ISCAM.docx")

# ── 0. Restaurar desde el respaldo más reciente pre-ingesta ────────────────
baks = sorted(glob.glob(os.path.join(ROOT, "ISCAM", "respaldos", "Metodologia.backup_*_pre-ingesta.docx")))
if not baks:
    sys.exit("no hay respaldo pre-ingesta del documento")
shutil.copy2(baks[-1], DOC)
print("restaurado desde:", os.path.basename(baks[-1]))

doc = docx.Document(DOC)
P = doc.paragraphs
p40, p57, p59, p68, p70, p72, p74 = P[40], P[57], P[59], P[68], P[70], P[72], P[74]


def insert_after(anchor, source, text):
    """Inserta un párrafo copiando el FORMATO de `source`, justo después de `anchor`."""
    new_p = copy.deepcopy(source._p)
    anchor._p.addnext(new_p)
    np = Paragraph(new_p, anchor._parent)
    for r in np.runs[1:]:
        r._r.getparent().remove(r._r)
    if np.runs:
        np.runs[0].text = text
    else:
        np.add_run(text)
    return np


# ── A) §2.2 — tabla de codificación ────────────────────────────────────────
insert_after(p40, p40,
    "La conversión de las respuestas categóricas a puntaje es explícita y auditable: la hoja CODIFICACION "
    "es la tabla de codificación del instrumento (una fila por opción de respuesta) y asigna a cada opción "
    "un valor de 0 a 100 —100 indica el mayor alineamiento con el EUDR y el Pacto Verde— o la marca «sin "
    "dato». Las preguntas de escala 0–4 se convierten ×25; las de porcentaje entran directas en 0–100; las "
    "de casillas múltiples puntúan por el número de componentes marcados y, en destinos de exportación, por "
    "la proporción de destinos de la Unión Europea. Las opciones que expresan desconocimiento o no "
    "aplicación («No sabe», «No verifica», «No se ha medido», «No aplica») se codifican como «sin dato» y no "
    "promedian: la ausencia de respuesta nunca se convierte en incumplimiento. La tabla es editable y marca "
    "con «[R]» los ítems cuyo sentido conviene ratificar con el autor antes del levantamiento.")

# ── B) §2.6 — capa de ingesta ─────────────────────────────────────────────
last = p59
bloque = [
    (p57, "2.6 Capa de ingesta: captura, codificación y normalización"),
    (p59, "El Instrumento 2 se carga por una sola vía. La plantilla incorpora tres hojas que separan "
          "entrada, reglas y cálculo, de modo que el índice se recalcule solo al pegar las respuestas:"),
    (p68, "CAPTURA — una fila por cooperativa (las 174 del marco muestral, con RUC y razón social "
          "precargados) y una columna por pregunta. Trae lista desplegable en las 63 preguntas de opción, "
          "validación numérica en las 10 de escala y porcentaje, y celda libre en las 4 de casillas "
          "múltiples. Es la única hoja donde se escribe."),
    (p68, "CODIFICACION — la tabla de codificación: opción de respuesta → puntaje 0-100 o «sin dato» "
          "(290 filas). Editable por diseño."),
    (p68, "NORMALIZADO — resultado automático: 0-100 por pregunta y el agregado de las 11 dimensiones "
          "(D1–D11), que alimenta directo al motor ISCAM_Score. No se edita."),
    (p59, "La unión con el resto de la plantilla es por RUC, la misma clave de todo el sistema. Las "
          "preguntas de casillas múltiples se separan con «;»; las ramas condicionales quedan vacías cuando "
          "no aplican; y una respuesta no codificada se señala con «⚠» para que no pase inadvertida. La "
          "dimensión D3 combina la encuesta con la fuente externa: si no hay respuesta, conserva el valor "
          "de la consulta OSCE/OEFA."),
]
for src, txt in bloque:
    last = insert_after(last, src, txt)

# ── C) §4 — número de hojas y reglas de carga ─────────────────────────────
p70.runs[0].text = p70.runs[0].text.replace("en las 7 áreas de la hoja", "en las 10 hojas del libro")
p72.runs[0].text = p72.runs[0].text.replace(
    "Reglas de integridad para quien cargue datos:",
    "Reglas de integridad para quien cargue datos (todo se escribe en la hoja CAPTURA):")

# ── D) §5 — la capa de ingesta deja de ser pendiente ──────────────────────
p74.runs[0].text = (
    "Capa de ingesta (implementada 2026-10-15). La tabla de codificación, la hoja de captura y la "
    "normalización están construidas y cableadas al motor (hojas CODIFICACION, CAPTURA y NORMALIZADO). "
    "Lo que resta es trabajo de campo: cargar las respuestas en CAPTURA. Cobertura de partida: D3 "
    "(legalidad OSCE/OEFA) 174/174; D1, D2, D4–D11 sin dato hasta aplicar el instrumento; D12 45/174; "
    "certificación 3/174. La consulta de legalidad OSCE/OEFA está completa para las 174 organizaciones "
    "(2026-10-15): 0 inhabilitaciones vigentes y 1 infracción ambiental OEFA.")

# ── E) T0 — nuevas áreas funcionales ──────────────────────────────────────
t0 = doc.tables[0]
for area, rol, cont in [
    ("Tabla de codificación (hoja CODIFICACION)", "Regla de conversión",
     "Las 290 opciones de respuesta con su puntaje 0-100 o «sin dato»; editable, marca «[R]» los ítems a ratificar"),
    ("Captura del instrumento (hoja CAPTURA)", "Entrada de datos",
     "174 cooperativas × 119 preguntas, con listas desplegables y validaciones; única hoja de escritura"),
    ("Normalización (hoja NORMALIZADO)", "Cálculo intermedio",
     "0-100 por pregunta y agregado de las 11 dimensiones (D1-D11); alimenta ISCAM_Score"),
]:
    row = copy.deepcopy(t0.rows[-1]._tr)
    t0._tbl.append(row)
    r = t0.rows[-1]
    for i, v in enumerate((area, rol, cont)):
        r.cells[i].text = v

for row in t0.rows:
    if row.cells[0].text.startswith("Scorecards"):
        row.cells[2].text = "ISA-Completo, semáforo, ranking; espejo de D2, D5 y D6 desde NORMALIZADO"

# ── F) T2 — normalización de categóricas ordinales ────────────────────────
t2 = doc.tables[2]
row = copy.deepcopy(t2.rows[1]._tr)
t2.rows[1]._tr.addnext(row)
r = t2.rows[1]
for i, v in enumerate(("Categórica ordinal (opción múltiple)", "Preguntas de opción (63)",
                       "Tabla de codificación (hoja CODIFICACION): 0-100 por opción; «sin dato» no promedia")):
    r.cells[i].text = v

doc.save(DOC)

# ── Verificación ──────────────────────────────────────────────────────────
doc = docx.Document(DOC)
print("párrafos:", len(doc.paragraphs), "| tablas:", len(doc.tables),
      "| T0 filas:", len(doc.tables[0].rows), "| T2 filas:", len(doc.tables[2].rows))
for i in range(58, 68):
    p = doc.paragraphs[i]
    r0 = p.runs[0]
    print(f"  P{i} [{p.style.name}] b={r0.bold} sz={r0.font.size} nm={r0.font.name}: {p.text.strip()[:78]}")

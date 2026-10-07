"""
Reconstruye RUC_FOB (exportaciones por RUC-anio-mes-aduana-agente-pais) en
Python -- es el nivel que faltaba: dua_canal_full.py depende de que RUC_FOB
ya exista, no lo genera. Este script sí lo genera desde cero.

Puerto directo de apps/identidad-fiscal/api/src/ingest/exportaciones-fob-connector.ts
+ exportaciones-fob-normalize.ts (repo Rastro/appsperu).

Endpoint: GET http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias
  ?accion=buscarListadoImpoExpo&CG_consulta=1&indOCE=false&vig001278=F
  &strMenu=-&CG_tipo=4&CG_Codigo={ruc}&CG_DNombre=&CG_Aduana=999
  &CG_Ano={anio-1992}&CG_Mes=00&CG_regimen=40

Sin captcha, confirmado en vivo por el conector original el 2026-09-19 y
otra vez desde este script (ver dua_canal_full.py, mismo dominio). El
parametro de anio NO es el anio real: hay que restarle 1992 -- offset
descubierto empiricamente por el conector original probando valores contra
resultados conocidos, no documentado por SUNAT/Aduanet en ningun lado.

El parseo replica exportaciones-fob-normalize.ts: cada fila de la tabla de
resultados trae, ademas del texto visible, los codigos reales (aduana,
agente, mes, anio, pais) embebidos en el atributo href del link "LISTAR"
(llamada JS a jsDetalleDUA(...)) -- se usan esos codigos en vez de parsear
el texto, porque no dependen del idioma/formato de fecha.

Uso:
  python ruc_fob_connector.py [--anios 2025,2026] [--limit N] [--resume-from N]
Lee el universo de RUC desde RUC_CONSULTAMASIVA (597 RUC) y escribe
ruc_fob.csv, con las mismas columnas que la pestana RUC_FOB. No toca el
xlsx.
"""
import argparse
import csv
import re
import sys
import time
from datetime import datetime, timezone
from urllib.request import urlopen, Request
from urllib.parse import urlencode

import openpyxl
from bs4 import BeautifulSoup

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
BASE_URL = "http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias"
ANIO_OFFSET = 1992
REGIMEN_EXPORTACION = "40"
DELAY_S = 0.4

JS_DETALLE_RE = re.compile(
    r'jsDetalleDUA\("(\d+)","(\d+)","(\d+)","(\d+)","(\d+)","(\d+)","([A-Za-z]{2})"\)'
)


def read_universo():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["RUC_CONSULTAMASIVA"]
    return [str(row[0].value) for row in ws.iter_rows(min_row=2, max_row=ws.max_row) if row[0].value]


def build_url(ruc, anio):
    params = {
        "accion": "buscarListadoImpoExpo",
        "CG_consulta": "1",
        "indOCE": "false",
        "vig001278": "F",
        "strMenu": "-",
        "CG_tipo": "4",
        "CG_Codigo": ruc,
        "CG_DNombre": "",
        "CG_Aduana": "999",
        "CG_Ano": str(anio - ANIO_OFFSET),
        "CG_Mes": "00",
        "CG_regimen": REGIMEN_EXPORTACION,
    }
    return f"{BASE_URL}?{urlencode(params)}"


def fetch(url, timeout=20):
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
    return raw.decode("latin1")


def clean_text(s):
    return re.sub(r"\s+", " ", (s or "").replace("\x00", "")).strip()


def parse_fob(s):
    try:
        return float(clean_text(s).replace(",", ""))
    except ValueError:
        return None


def parse_exportaciones_html(html):
    soup = BeautifulSoup(html, "html.parser")
    rows = []
    for tr in soup.select("tr.bg"):
        tds = tr.find_all("td")
        if len(tds) < 7:
            continue
        a = tds[0].find("a")
        href = a.get("href", "") if a else ""
        m = JS_DETALLE_RE.search(href)
        if not m:
            continue
        _, tip_nro_doc, cod_aduana, cod_agente, cod_mes, cod_anio, cod_pais = m.groups()
        ruc = tip_nro_doc[-11:]

        agente_nombre = clean_text(tds[3].get_text()) or None
        aduana_nombre = clean_text(tds[4].get_text()) or None
        pais_nombre = clean_text(tds[5].get_text()) or None
        fob_usd = parse_fob(tds[6].get_text())
        if fob_usd is None:
            continue

        rows.append({
            "ruc": ruc,
            "anio": int(cod_anio),
            "mes": int(cod_mes),
            "aduana_codigo": cod_aduana,
            "aduana_nombre": aduana_nombre,
            "agente_codigo": cod_agente,
            "agente_nombre": agente_nombre,
            "pais_codigo": cod_pais.upper(),
            "pais_nombre": pais_nombre,
            "fob_usd": fob_usd,
        })
    return rows


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--anios", default="2025,2026", help="Anios a consultar, separados por coma")
    ap.add_argument("--limit", type=int, default=None, help="Procesar solo los primeros N RUC (pruebas)")
    ap.add_argument("--resume-from", type=int, default=0)
    ap.add_argument("--out", default="ruc_fob.csv")
    args = ap.parse_args()

    anios = [int(a.strip()) for a in args.anios.split(",")]
    rucs = read_universo()
    print(f"universo: {len(rucs)} RUC, anios: {anios}")

    if args.limit:
        rucs = rucs[args.resume_from: args.resume_from + args.limit]
    else:
        rucs = rucs[args.resume_from:]

    fieldnames = ["ruc", "anio", "mes", "aduana_codigo", "aduana_nombre", "agente_codigo",
                  "agente_nombre", "pais_codigo", "pais_nombre", "fob_usd", "fecha_consulta"]
    mode = "a" if args.resume_from else "w"
    write_header = not (args.resume_from and mode == "a")

    n_ok, n_con_datos, n_err = 0, 0, 0
    t0 = time.time()

    with open(args.out, mode, newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        if write_header:
            w.writeheader()

        total = len(rucs) * len(anios)
        i = 0
        for ruc in rucs:
            ruc_tuvo_datos = False
            for anio in anios:
                i += 1
                fecha_consulta = datetime.now(timezone.utc).isoformat()
                try:
                    html = fetch(build_url(ruc, anio))
                    rows = parse_exportaciones_html(html)
                    if rows:
                        ruc_tuvo_datos = True
                    for r in rows:
                        r["fecha_consulta"] = fecha_consulta
                        w.writerow(r)
                    n_ok += 1
                except Exception as e:
                    n_err += 1
                    print(f"[{i}/{total}] ERROR ruc={ruc} anio={anio}: {e}", flush=True)

                if i % 50 == 0 or i == total:
                    elapsed = time.time() - t0
                    print(f"[{i}/{total}] ok={n_ok} err={n_err} ({elapsed:.0f}s, {elapsed/i:.2f}s/consulta)", flush=True)

                time.sleep(DELAY_S)
            if ruc_tuvo_datos:
                n_con_datos += 1

    print(f"Listo. consultas_ok={n_ok} err={n_err} RUC_con_exportaciones={n_con_datos}. CSV: {args.out}")


if __name__ == "__main__":
    sys.exit(main())

"""
Ingesta completa de numero de DUA + canal de control (verde/rojo/naranja) por
cada operacion de RUC_FOB, en un solo script (antes eran dos pasadas
manuales: dua_canal_connector.py + dua_canal_retry_2024.py).

Endpoint: GET http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias
  ?accion=buscarListadoDuas&CG_consulta=1&CG_regimen=40&CG_ruc=4{ruc}
  &CG_aduana={aduana_codigo}&CG_agente={agente_codigo}&CG_mes={mes}
  &CG_anio={anio}&CG_pais={pais_codigo}
Sin captcha, confirmado en vivo 24/09/2026 (mismo dominio que
exportaciones-fob-connector.ts del repo Rastro). El detalle de partida
arancelaria de cada DUA (servlet SgCDUI2) SI pide captcha -- no se
automatiza, ver hallazgo del 23/09/2026.

Dos niveles en una sola corrida:
  Nivel 1 -- filas de RUC_FOB que ya traen agente_codigo (la mayoria: la
             consulta original 2025-2026, sin captcha).
  Nivel 2 -- filas que NO traen agente_codigo (el archivo historico 2024
             aportado por el usuario solo trajo agente_nombre). Se resuelve
             el codigo cruzando el nombre normalizado contra las filas del
             Nivel 1 que SI lo tienen; los nombres que no encuentran match
             (agentes que solo aparecen en 2024) quedan sin DUA, reportados
             al final -- no se inventa el codigo.

Uso:
  python dua_canal_full.py [--limit N] [--resume-from N] [--out CSV]
Lee "BASE CONSOLIDADA.xlsx" (pestana RUC_FOB) y escribe UN SOLO CSV final
(dua_canal_final.csv por defecto) con ambos niveles ya combinados -- no hace
falta ningun paso manual de union despues. No toca el xlsx (eso lo hace
load_dua_canal.py, aparte).
"""
import argparse
import csv
import re
import sys
import time
from datetime import datetime, timezone
from urllib.request import urlopen, Request
from urllib.parse import urlencode

import openpyxl
from bs4 import BeautifulSoup

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
BASE_URL = "http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias"
DELAY_S = 0.4


def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().upper())


def read_ruc_fob_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["RUC_FOB"]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        vals = [c.value for c in row]
        if vals[0] is None:
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


def resolve_missing_agente_codigo(rows):
    """Nivel 2: para filas sin agente_codigo, lo resuelve por nombre
    normalizado contra las filas que si lo tienen. Devuelve
    (filas_listas_para_consultar, nombres_sin_resolver)."""
    lookup = {}
    for r in rows:
        if r.get("agente_codigo"):
            lookup.setdefault(norm(r["agente_nombre"]), set()).add(r["agente_codigo"])

    listas = []
    sin_resolver = []
    for r in rows:
        if r.get("agente_codigo"):
            listas.append(r)
            continue
        codes = lookup.get(norm(r.get("agente_nombre")))
        if codes and len(codes) == 1:
            r2 = dict(r)
            r2["agente_codigo"] = next(iter(codes))
            listas.append(r2)
        else:
            sin_resolver.append(r)
    return listas, sin_resolver


def build_url(ruc, anio, mes, aduana_codigo, agente_codigo, pais_codigo):
    params = {
        "accion": "buscarListadoDuas",
        "CG_consulta": "1",
        "CG_regimen": "40",
        "CG_ruc": f"4{ruc}",
        "CG_aduana": str(aduana_codigo),
        "CG_agente": str(agente_codigo),
        "CG_mes": str(mes),
        "CG_anio": str(anio),
        "CG_pais": pais_codigo,
    }
    return f"{BASE_URL}?{urlencode(params)}"


def fetch(url, timeout=20):
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
    return raw.decode("latin1")


def parse_duas(html):
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for tr in soup.select("tr.bg"):
        tds = tr.find_all("td")
        if len(tds) < 5:
            continue
        aduana = tds[0].get_text(strip=True)
        anio = tds[1].get_text(strip=True)
        a = tds[2].find("a")
        if a is None:
            continue
        dua_numero = a.get_text(strip=True)
        href = a.get("href", "")
        fecha_embarque = tds[3].get_text(strip=True)
        canal = tds[4].get_text(strip=True).upper()
        if not aduana.isdigit() or not anio.isdigit():
            continue
        out.append({
            "aduana_codigo": aduana, "anio_dua": anio, "dua_numero": dua_numero,
            "fecha_embarque": fecha_embarque, "canal": canal, "url_detalle": href,
        })
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None, help="Procesar solo las primeras N filas (pruebas)")
    ap.add_argument("--resume-from", type=int, default=0, help="Saltar las primeras N filas a consultar (para reanudar)")
    ap.add_argument("--out", default="dua_canal_final.csv")
    args = ap.parse_args()

    todas = read_ruc_fob_rows()
    print(f"filas RUC_FOB: {len(todas)}")

    a_consultar, sin_resolver = resolve_missing_agente_codigo(todas)
    print(f"Nivel 1 (ya tenian agente_codigo): {sum(1 for r in todas if r.get('agente_codigo'))}")
    print(f"Nivel 2 (resueltas por nombre de agente): {len(a_consultar) - sum(1 for r in todas if r.get('agente_codigo'))}")
    if sin_resolver:
        nombres = sorted(set(r.get("agente_nombre") for r in sin_resolver))
        print(f"Sin resolver ({len(sin_resolver)} filas, {len(nombres)} agentes) -- no se consultan, quedan sin DUA:")
        for n in nombres:
            print(f"  - {n}")

    if args.limit:
        a_consultar = a_consultar[args.resume_from: args.resume_from + args.limit]
    else:
        a_consultar = a_consultar[args.resume_from:]

    fieldnames = [
        "ruc", "anio", "mes", "aduana_codigo", "agente_codigo", "pais_codigo",
        "dua_numero", "fecha_embarque", "canal", "url_detalle", "fecha_consulta",
    ]
    mode = "a" if args.resume_from else "w"
    write_header = not (args.resume_from and mode == "a")
    n_ok, n_sin_dua, n_err = 0, 0, 0
    t0 = time.time()

    with open(args.out, mode, newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        if write_header:
            w.writeheader()

        for i, r in enumerate(a_consultar, start=1):
            url = build_url(r["ruc"], r["anio"], r["mes"], r["aduana_codigo"], r["agente_codigo"], r["pais_codigo"])
            fecha_consulta = datetime.now(timezone.utc).isoformat()
            try:
                html = fetch(url)
                duas = parse_duas(html)
                if not duas:
                    n_sin_dua += 1
                for d in duas:
                    w.writerow({
                        "ruc": r["ruc"], "anio": r["anio"], "mes": r["mes"],
                        "aduana_codigo": r["aduana_codigo"], "agente_codigo": r["agente_codigo"],
                        "pais_codigo": r["pais_codigo"], "dua_numero": d["dua_numero"],
                        "fecha_embarque": d["fecha_embarque"], "canal": d["canal"],
                        "url_detalle": d["url_detalle"], "fecha_consulta": fecha_consulta,
                    })
                n_ok += 1
            except Exception as e:
                n_err += 1
                print(f"[{i}/{len(a_consultar)}] ERROR ruc={r['ruc']} anio={r['anio']} mes={r['mes']}: {e}", flush=True)

            if i % 50 == 0 or i == len(a_consultar):
                elapsed = time.time() - t0
                print(f"[{i}/{len(a_consultar)}] ok={n_ok} sin_dua={n_sin_dua} err={n_err} "
                      f"({elapsed:.0f}s, {elapsed/i:.2f}s/fila)", flush=True)

            time.sleep(DELAY_S)

    print(f"Listo. ok={n_ok} sin_dua={n_sin_dua} err={n_err}. CSV final: {args.out}")
    print(f"Filas de RUC_FOB sin DUA por agente no resoluble: {len(sin_resolver)}")


if __name__ == "__main__":
    sys.exit(main())

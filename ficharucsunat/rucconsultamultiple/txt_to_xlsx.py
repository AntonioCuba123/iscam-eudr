import sys
import io
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

def read_delim(src: str) -> pd.DataFrame:
    with open(src, "rb") as f:
        raw = f.read()
    for enc in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return pd.read_csv(io.BytesIO(raw), sep="|", dtype=str, encoding=enc, keep_default_na=False, na_filter=False)
        except UnicodeDecodeError:
            continue
    return pd.read_csv(io.BytesIO(raw), sep="|", dtype=str, encoding="latin-1", keep_default_na=False, na_filter=False)


def txt_to_xlsx(src: str, dst: str):
    df = read_delim(src)
    df.columns = [c.strip() for c in df.columns]
    df = df.loc[:, [c for c in df.columns if c != "" and not str(c).startswith("Unnamed")]]
    df = df.replace({"": pd.NA})
    if "NumeroRuc" in df.columns:
        df["NumeroRuc"] = df["NumeroRuc"].astype("string")

    with pd.ExcelWriter(dst, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="RUC")
        ws = writer.sheets["RUC"]

        header_fill = PatternFill("solid", fgColor="1F3864")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 32

        for idx, col in enumerate(df.columns, start=1):
            letter = get_column_letter(idx)
            if col in ("NumeroRuc",):
                width = 14
            elif col in ("Nombre ó RazonSocial", "Direccion", "Actividad de Comercio Exterior"):
                width = 52
            elif "CIIU" in col or "Actividad" in col or "Comercio" in col:
                width = 34
            else:
                width = max(14, min(26, len(col) + 4))
            ws.column_dimensions[letter].width = width
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=False)

        for cell in ws["A"][1:]:
            cell.alignment = Alignment(horizontal="center", vertical="top")
            cell.number_format = "@"

    print(f"filas: {len(df)} | columnas: {len(df.columns)} | out: {dst}")

if __name__ == "__main__":
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else src.rsplit(".", 1)[0] + ".xlsx"
    txt_to_xlsx(src, dst)

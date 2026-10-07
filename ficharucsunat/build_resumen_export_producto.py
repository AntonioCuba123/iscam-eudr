"""
Construye la pestana RESUMEN_EXPORT_VENTAS_PRODUCTO: exportacion sobre ventas,
agrupado por producto_dominante.

Dos fuentes de "exportacion" distintas -- se muestran con nombres separados,
NUNCA mezcladas en una sola cifra:

1. exportacion_b2_pen -- autoreportado por las organizaciones (ventas_exportacion
   de HECHO_PROGRAMA_ANUAL / b2). Misma moneda (PEN) y misma fuente que
   ventas_totales, asi que el ratio exportacion/ventas SI es valido de calcular
   directo.
2. exportacion_sunat_usd_total -- fob_total_usd de DIM_ORGANIZACION (verificado
   contra SUNAT-Aduanas via RUC_FOB). En USD, todos los anios acumulados
   (2024-2026), y solo cubre a los RUC que SUNAT registra como exportadores
   (66 de 597). Se deja como cifra de referencia aparte, SIN dividir entre
   ventas_totales: se probo convertir a soles con un tipo de cambio de
   referencia y comparar contra ventas de un solo anio, pero el resultado da
   ratios sobre 100% para CAFE y CACAO -- no es un error de calculo, es que
   ventas_totales (b2) parece medir algo mas angosto que el total comercial
   real de la organizacion (ver casos verificados: COOPERATIVA ASPRAOSRA
   reporta S/782 mil de ventas totales 2025 pero SUNAT verifica US$6.78M en
   exportaciones). Se quito la columna estimada de la pestana para no sugerir
   una precision que no existe; el dato USD queda solo como referencia.

Producto: producto_dominante de DIM_ORGANIZACION (fuente b2, respaldo b1). Los
RUC que estan en b3 (programa de cafe PAC) pero sin producto_dominante propio
se completan con "CAFE" -- b3 es integramente un programa cafetalero, no hay
ambiguedad de producto para esos casos.

Snapshot calculado con Python, no formulas de Excel en vivo (mismo motivo que
el resto de correcciones de hoy: las formulas de array dinamico de este libro
ya mostraron ser fragiles ante cambios de tamano de los datos). Reproducible:
volver a correr este script despues de actualizar fuentes.
"""
from datetime import datetime, timezone

import openpyxl
from openpyxl.styles import Font, PatternFill

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
SHEET_NAME = "RESUMEN_EXPORT_VENTAS_PRODUCTO"


def sheet_rows(wb, name, headerrow=1):
    ws = wb[name]
    headers = [c.value for c in ws[headerrow]]
    rows = []
    for row in ws.iter_rows(min_row=headerrow + 1, max_row=ws.max_row):
        vals = [c.value for c in row]
        if vals[0] is None:
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

    dim = {r["ruc"]: r for r in sheet_rows(wb, "DIM_ORGANIZACION")}
    hpa = sheet_rows(wb, "HECHO_PROGRAMA_ANUAL")
    b3_rucs = set(str(r["RUC"]) for r in sheet_rows(wb, "b3"))

    # Salvaguarda: si DIM_ORGANIZACION se guardo con openpyxl despues de la
    # ultima recalculacion en Excel, la fila-ancla de sus formulas de array
    # dinamico (fila 2) puede perder su valor cacheado hasta que Excel
    # recalcule -- eso deja un RUC entero fuera del dict. Se detecta contra
    # el universo real (RUC_CONSULTAMASIVA, hoja de datos planos, inmune a
    # este problema) y se reconstruye ese RUC desde las fuentes primarias.
    universo = set(str(r["NumeroRuc"]) for r in sheet_rows(wb, "RUC_CONSULTAMASIVA"))
    faltantes = universo - set(dim.keys())
    if faltantes:
        print(f"AVISO: {len(faltantes)} RUC ausentes de DIM_ORGANIZACION por cache de formula "
              f"(recalcular en Excel para evitar esto). Reconstruyendo desde fuentes primarias: {faltantes}")
        # b2 no tiene headers de texto limpio en todas las columnas -- se lee por
        # posicion: RUC = columna 30, PRODUCTO DOM = columna 2.
        ws_b2 = wb["b2"]
        prod_por_ruc = {}
        for row in ws_b2.iter_rows(min_row=3, max_row=ws_b2.max_row):
            ruc_val = row[29].value  # columna 30 = RUC
            prod_val = row[1].value  # columna 2 = PRODUCTO DOM
            if ruc_val and str(ruc_val) in faltantes and prod_val:
                prod_por_ruc[str(ruc_val)] = prod_val

        fob_por_ruc = {}
        for r in sheet_rows(wb, "HECHO_EXPORTACION"):
            ruc = str(r["ruc"])
            if ruc in faltantes:
                fob_por_ruc[ruc] = fob_por_ruc.get(ruc, 0) + (r.get("fob_usd") or 0)

        ventas_por_ruc_anio = {}
        for r in sheet_rows(wb, "HECHO_PROGRAMA_ANUAL"):
            ruc = str(r["ruc"])
            if ruc in faltantes and r["anio"] in (2024, 2025):
                key = (ruc, r["anio"])
                ventas_por_ruc_anio[key] = ventas_por_ruc_anio.get(key, 0) + (r.get("ventas_totales") or 0)

        for ruc in faltantes:
            dim[ruc] = {
                "producto_dominante": prod_por_ruc.get(ruc),
                "ventas_totales_2024": ventas_por_ruc_anio.get((ruc, 2024)),
                "ventas_totales_2025": ventas_por_ruc_anio.get((ruc, 2025)),
                "fob_total_usd": fob_por_ruc.get(ruc, 0),
            }

    # exportacion / ventas autoreportadas (b2) por ruc y anio
    export_b2 = {}  # (ruc, anio) -> ventas_exportacion
    for r in hpa:
        if r["anio"] in (2024, 2025):
            export_b2[(r["ruc"], r["anio"])] = (r.get("ventas_exportacion") or 0)

    # producto por ruc, con respaldo CAFE para RUC de b3 sin producto_dominante
    producto_por_ruc = {}
    sin_producto_pero_b3 = 0
    sin_producto_total = 0
    for ruc, r in dim.items():
        p = (r.get("producto_dominante") or "").strip()
        if not p:
            if ruc in b3_rucs:
                p = "CAFÉ"
                sin_producto_pero_b3 += 1
            else:
                p = "SIN PRODUCTO REGISTRADO"
                sin_producto_total += 1
        producto_por_ruc[ruc] = p

    print(f"RUC sin producto_dominante resueltos via b3 (CAFE): {sin_producto_pero_b3}")
    print(f"RUC sin producto_dominante y sin b3 (quedan 'SIN PRODUCTO REGISTRADO'): {sin_producto_total}")

    # agregacion por producto
    grupos = {}
    for ruc, r in dim.items():
        p = producto_por_ruc[ruc]
        g = grupos.setdefault(p, {
            "n_organizaciones": 0,
            "ventas_totales_2024_pen": 0.0,
            "exportacion_b2_2024_pen": 0.0,
            "ventas_totales_2025_pen": 0.0,
            "exportacion_b2_2025_pen": 0.0,
            "exportacion_sunat_usd_total": 0.0,
            "n_organizaciones_con_fob_sunat": 0,
        })
        g["n_organizaciones"] += 1
        g["ventas_totales_2024_pen"] += (r.get("ventas_totales_2024") or 0)
        g["exportacion_b2_2024_pen"] += export_b2.get((ruc, 2024), 0)
        g["ventas_totales_2025_pen"] += (r.get("ventas_totales_2025") or 0)
        g["exportacion_b2_2025_pen"] += export_b2.get((ruc, 2025), 0)
        fob = r.get("fob_total_usd") or 0
        g["exportacion_sunat_usd_total"] += fob
        if fob:
            g["n_organizaciones_con_fob_sunat"] += 1

    # armar filas ordenadas por ventas_totales_2025 desc, con fila TOTAL al final
    filas = []
    for producto, g in grupos.items():
        pct_2024 = (g["exportacion_b2_2024_pen"] / g["ventas_totales_2024_pen"]
                    if g["ventas_totales_2024_pen"] else None)
        pct_2025 = (g["exportacion_b2_2025_pen"] / g["ventas_totales_2025_pen"]
                    if g["ventas_totales_2025_pen"] else None)
        filas.append((producto, g, pct_2024, pct_2025))
    filas.sort(key=lambda x: x[1]["ventas_totales_2025_pen"], reverse=True)

    total = {
        "n_organizaciones": sum(g["n_organizaciones"] for _, g, _, _ in filas),
        "ventas_totales_2024_pen": sum(g["ventas_totales_2024_pen"] for _, g, _, _ in filas),
        "exportacion_b2_2024_pen": sum(g["exportacion_b2_2024_pen"] for _, g, _, _ in filas),
        "ventas_totales_2025_pen": sum(g["ventas_totales_2025_pen"] for _, g, _, _ in filas),
        "exportacion_b2_2025_pen": sum(g["exportacion_b2_2025_pen"] for _, g, _, _ in filas),
        "exportacion_sunat_usd_total": sum(g["exportacion_sunat_usd_total"] for _, g, _, _ in filas),
        "n_organizaciones_con_fob_sunat": sum(g["n_organizaciones_con_fob_sunat"] for _, g, _, _ in filas),
    }
    total_pct_2024 = (total["exportacion_b2_2024_pen"] / total["ventas_totales_2024_pen"]
                       if total["ventas_totales_2024_pen"] else None)
    total_pct_2025 = (total["exportacion_b2_2025_pen"] / total["ventas_totales_2025_pen"]
                       if total["ventas_totales_2025_pen"] else None)

    # escribir la pestana
    wbw = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    if SHEET_NAME in wbw.sheetnames:
        del wbw[SHEET_NAME]
    ws = wbw.create_sheet(SHEET_NAME)

    ws["A1"] = (
        "Exportacion sobre ventas, por producto. Dos fuentes de exportacion NO comparables entre si "
        "(monedas y coberturas distintas) -- ver columnas separadas. exportacion_b2_*_pen: "
        "autoreportado por las organizaciones (HECHO_PROGRAMA_ANUAL/b2), misma moneda que "
        "ventas_totales, ratio valido. exportacion_sunat_usd_total: verificado SUNAT-Aduanas "
        "(RUC_FOB), TODOS los anios acumulados (2024-2026), en USD -- cifra de referencia, NO "
        "dividir entre ventas_totales (se probo con tipo de cambio fijo y da ratios sobre 100% para "
        "CAFE/CACAO: ventas_totales de b2 parece medir algo mas angosto que el total comercial real "
        "de la organizacion, no un error de calculo -- ver casos verificados en el docstring del "
        "script que genera esta pestana). Producto = producto_dominante de DIM_ORGANIZACION (fuente "
        "b2, respaldo b1); RUC de b3 sin producto propio se marcan CAFE (b3 es integramente el "
        "programa cafetalero PAC)."
    )
    ws.merge_cells("A1:K1")
    ws["A1"].font = Font(italic=True, size=9)
    ws.row_dimensions[1].height = 60
    ws["A1"].alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")

    headers = [
        "producto_dominante", "n_organizaciones",
        "ventas_totales_2024_pen", "exportacion_b2_2024_pen", "pct_exportacion_ventas_2024_b2",
        "ventas_totales_2025_pen", "exportacion_b2_2025_pen", "pct_exportacion_ventas_2025_b2",
        "exportacion_sunat_usd_total", "n_organizaciones_con_fob_sunat",
        "fecha_calculo",
    ]
    ws.append([None] * 0)  # fila 2 en blanco (separacion visual)
    header_row = 3
    for j, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=j, value=h)

    fecha_calculo = datetime.now(timezone.utc).isoformat()
    r_idx = header_row + 1
    for producto, g, pct_2024, pct_2025 in filas:
        ws.cell(row=r_idx, column=1, value=producto)
        ws.cell(row=r_idx, column=2, value=g["n_organizaciones"])
        ws.cell(row=r_idx, column=3, value=round(g["ventas_totales_2024_pen"], 2))
        ws.cell(row=r_idx, column=4, value=round(g["exportacion_b2_2024_pen"], 2))
        ws.cell(row=r_idx, column=5, value=round(pct_2024, 4) if pct_2024 is not None else None)
        ws.cell(row=r_idx, column=6, value=round(g["ventas_totales_2025_pen"], 2))
        ws.cell(row=r_idx, column=7, value=round(g["exportacion_b2_2025_pen"], 2))
        ws.cell(row=r_idx, column=8, value=round(pct_2025, 4) if pct_2025 is not None else None)
        ws.cell(row=r_idx, column=9, value=round(g["exportacion_sunat_usd_total"], 2))
        ws.cell(row=r_idx, column=10, value=g["n_organizaciones_con_fob_sunat"])
        ws.cell(row=r_idx, column=11, value=fecha_calculo)
        r_idx += 1

    # fila TOTAL GENERAL, resaltada
    total_row = r_idx
    ws.cell(row=total_row, column=1, value="TOTAL GENERAL")
    ws.cell(row=total_row, column=2, value=total["n_organizaciones"])
    ws.cell(row=total_row, column=3, value=round(total["ventas_totales_2024_pen"], 2))
    ws.cell(row=total_row, column=4, value=round(total["exportacion_b2_2024_pen"], 2))
    ws.cell(row=total_row, column=5, value=round(total_pct_2024, 4) if total_pct_2024 is not None else None)
    ws.cell(row=total_row, column=6, value=round(total["ventas_totales_2025_pen"], 2))
    ws.cell(row=total_row, column=7, value=round(total["exportacion_b2_2025_pen"], 2))
    ws.cell(row=total_row, column=8, value=round(total_pct_2025, 4) if total_pct_2025 is not None else None)
    ws.cell(row=total_row, column=9, value=round(total["exportacion_sunat_usd_total"], 2))
    ws.cell(row=total_row, column=10, value=total["n_organizaciones_con_fob_sunat"])
    ws.cell(row=total_row, column=11, value=fecha_calculo)

    yellow = PatternFill(fill_type="solid", fgColor="FFFFF2CC")
    for c in ws[total_row]:
        c.fill = yellow
        c.font = Font(bold=True)

    blue = PatternFill(fill_type="solid", fgColor="FFD9E1F2")
    for c in ws[header_row]:
        c.font = Font(bold=True)
        c.fill = blue

    # formato porcentaje en las columnas de ratio
    for col_pct in (5, 8):
        for row in ws.iter_rows(min_row=header_row + 1, max_row=total_row, min_col=col_pct, max_col=col_pct):
            for cell in row:
                cell.number_format = "0.0%"

    ws.freeze_panes = f"A{header_row + 1}"
    for col, width in zip("ABCDEFGHIJK", [26, 15, 18, 18, 14, 18, 18, 14, 18, 14, 24]):
        ws.column_dimensions[col].width = width

    wbw.save(XLSX_PATH)
    print(f"Guardado OK: pestana {SHEET_NAME}, {len(filas)} productos + 1 fila TOTAL")


if __name__ == "__main__":
    main()

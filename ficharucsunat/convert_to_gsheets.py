"""
Convierte BASE CONSOLIDADA_GSheets.xlsx para que sea abrible en Google Sheets.

Dos problemas de compatibilidad, no uno:
1. ANCHORARRAY (el operador de "spill" # de Excel 365) no existe en Sheets.
2. Excel guarda las funciones nuevas (FILTER, SORT, XLOOKUP, LET, UNIQUE, LAMBDA,
   MAP, VSTACK, SEQUENCE, TAKE, BYROW, HSTACK, MINIFS, MAXIFS, TOCOL) con el
   prefijo interno "_xlfn." (y "_xlfn._xlws." para FILTER/SORT, que chocan con
   nombres viejos). El importador de Sheets NO los reconoce con ese prefijo
   literal (da #NAME?, función desconocida ".SORT", etc.) — hay que dejarlos
   como el nombre de función plano (SORT, FILTER, XLOOKUP...) para que Sheets
   los entienda.

Estrategia:
1. ANCHORARRAY(ref) se reemplaza por el rango explicito exacto que hoy
   devuelve, buscando la formula de array ancla en esa misma celda (misma
   hoja, mismo ref) y leyendo su .ref (p.ej. A2:A598). Traduccion 1:1 del
   estado actual — reproduce exactamente lo que Excel calcula hoy, sin
   inventar buffers.
2. Se le quita el prefijo "_xlfn." (y "_xlfn._xlws.") a TODAS las formulas
   del libro, no solo a las que tenian ANCHORARRAY.

Limitacion: al ser rangos explicitos (no "spill" dinamico), si despues crecen
las fuentes (RUC_CONSULTAMASIVA, RUC_OECE, etc.) esta copia para Sheets NO se
va a expandir sola. Volver a correr este script sobre una copia fresca de
BASE CONSOLIDADA.xlsx regenera BASE CONSOLIDADA_GSheets.xlsx con los rangos
actualizados.

Uso:
  python convert_to_gsheets.py
Lee y sobreescribe BASE CONSOLIDADA_GSheets.xlsx (asume que ya es una copia de
BASE CONSOLIDADA.xlsx en la misma carpeta).
"""
import re
import sys

import openpyxl
from openpyxl.worksheet.formula import ArrayFormula

PATH = "BASE CONSOLIDADA_GSheets.xlsx"

CELL_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+)$")
REF_RE = re.compile(r"^\$?([A-Za-z]{1,3})\$?(\d+):\$?([A-Za-z]{1,3})\$?(\d+)$")
ANCHOR_CALL_RE = re.compile(r"_xlfn\.ANCHORARRAY\(([^()]*)\)")
PREFIX_RE = re.compile(r"_xlfn\._xlws\.|_xlfn\.")


def strip_xlfn_prefixes(text):
    return PREFIX_RE.sub("", text)


# MAP(...) confirmado roto en Sheets del usuario (BYROW, LAMBDA, TAKE, TOCOL,
# SEQUENCE ya confirmados OK). Reescritas a mano las 5 formulas que usan MAP:
# - caso de un solo array: MAP(range, LAMBDA(x, body)) -> BYROW(range, LAMBDA(x, body))
#   (con un solo array, cada "fila" que entrega BYROW es 1x1 y se comporta igual
#   que el escalar que entregaba MAP)
# - caso de varios arrays: se arman en una sola matriz con CHOOSE({1,2,...}, r1, r2, ...)
#   (viejo truco de Excel/Sheets para "apilar" columnas sin HSTACK) y se
#   desarma adentro del LAMBDA con INDEX(fila,1,i)
MAP_REPLACEMENTS = {
    ("DIM_ORGANIZACION", "AE2"): (
        "=BYROW($A$2:$A$598,LAMBDA(_xlpm.r,LET(_xlpm.x,"
        "FILTER(HECHO_EXPORTACION!$D$2:$D$959,HECHO_EXPORTACION!$A$2:$A$959=_xlpm.r,\"\"),"
        "_xlpm.m,MAX(IFERROR(--_xlpm.x,0)),"
        "IF(_xlpm.m=0,\"\",LEFT(_xlpm.m,4)&\"-\"&RIGHT(_xlpm.m,2)))))"
    ),
    ("DIM_ORGANIZACION", "AG2"): (
        "=BYROW($A$2:$A$598,LAMBDA(_xlpm.r,"
        "IF(COUNTIF(HECHO_EXPORTACION!$A$2:$A$959,_xlpm.r)=0,\"\","
        "ROWS(UNIQUE(FILTER(HECHO_EXPORTACION!$J$2:$J$959,HECHO_EXPORTACION!$A$2:$A$959=_xlpm.r))))))"
    ),
    ("REL_PERSONA_CARGO", "H2"): (
        "=BYROW(CHOOSE({1,2,3,4,5},$A$2:$A$816,$D$2:$D$816,$F$2:$F$816,$G$2:$G$816,$I$2:$I$816),"
        "LAMBDA(_xlpm.row,LET("
        "_xlpm.a,INDEX(_xlpm.row,1,1),_xlpm.d,INDEX(_xlpm.row,1,2),_xlpm.f,INDEX(_xlpm.row,1,3),"
        "_xlpm.g,INDEX(_xlpm.row,1,4),_xlpm.v,INDEX(_xlpm.row,1,5),"
        "_xlpm.n,MINIFS($G$2:$G$816,$A$2:$A$816,_xlpm.a,$D$2:$D$816,_xlpm.d,$F$2:$F$816,_xlpm.f,"
        "$G$2:$G$816,\">\"&N(_xlpm.g)),"
        "IF(OR(_xlpm.v=\"SI\",_xlpm.n=0),\"\",_xlpm.n))))"
    ),
    ("REL_PERSONA_CARGO", "I2"): (
        "=BYROW(CHOOSE({1,2,3,4},$A$2:$A$816,$D$2:$D$816,$F$2:$F$816,$G$2:$G$816),"
        "LAMBDA(_xlpm.row,LET("
        "_xlpm.a,INDEX(_xlpm.row,1,1),_xlpm.d,INDEX(_xlpm.row,1,2),"
        "_xlpm.f,INDEX(_xlpm.row,1,3),_xlpm.g,INDEX(_xlpm.row,1,4),"
        "IF(N(_xlpm.g)>=MAXIFS($G$2:$G$816,$A$2:$A$816,_xlpm.a,$D$2:$D$816,_xlpm.d,$F$2:$F$816,_xlpm.f),"
        "\"SI\",\"NO\"))))"
    ),
    ("DIM_PERSONA", "C2"): (
        "=BYROW($A$2:$A$713,LAMBDA(_xlpm.x,"
        "ROWS(UNIQUE(FILTER(REL_PERSONA_CARGO!$A$2:$A$816,REL_PERSONA_CARGO!$B$2:$B$816=_xlpm.x)))))"
    ),
}


def apply_map_replacements(wb):
    applied = 0
    for (sheetname, coord), new_text in MAP_REPLACEMENTS.items():
        ws = wb[sheetname]
        cell = ws[coord]
        old = cell.value
        if not isinstance(old, ArrayFormula):
            print(f"AVISO: {sheetname}!{coord} ya no es ArrayFormula, se omite el reemplazo de MAP")
            continue
        if "MAP(" not in strip_xlfn_prefixes(old.text):
            continue  # ya convertido en una corrida anterior
        cell.value = ArrayFormula(ref=old.ref, text=new_text)
        applied += 1
    return applied


def build_anchor_index(wb):
    """Mapea (nombre_hoja, 'COL#') -> ref completo (p.ej. 'A2:A598') de cada
    formula de array ancla en el libro."""
    index = {}
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    index[(sheetname, c.coordinate)] = v.ref
    return index


def resolve_anchor_arg(arg, current_sheet, index):
    arg = arg.strip()
    if "!" in arg:
        sheet_part, cellref = arg.split("!", 1)
        sheet_part = sheet_part.strip("'")
    else:
        sheet_part = None
        cellref = arg

    m = CELL_RE.match(cellref)
    if not m:
        raise ValueError(f"No pude parsear referencia de celda: {arg!r}")
    col, row = m.group(1).upper(), m.group(2)

    target_sheet = sheet_part or current_sheet
    key = (target_sheet, f"{col}{row}")
    if key not in index:
        raise KeyError(f"No encontre formula de array ancla en {key} (arg original: {arg!r})")

    ref = index[key]
    rm = REF_RE.match(ref)
    if not rm:
        raise ValueError(f"Ref inesperado {ref!r} para {key}")
    end_row = rm.group(4)

    prefix = f"{sheet_part}!" if sheet_part else ""
    return f"{prefix}${col}${row}:${col}${end_row}"


def convert_formula_text(text, current_sheet, index):
    def _sub(m):
        arg = m.group(1)
        return resolve_anchor_arg(arg, current_sheet, index)

    return ANCHOR_CALL_RE.sub(_sub, text)


def main():
    wb = openpyxl.load_workbook(PATH, data_only=False)
    index = build_anchor_index(wb)
    print(f"Formulas de array ancla indexadas: {len(index)}")

    changed_anchor = 0
    changed_prefix = 0
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, ArrayFormula):
                    text = v.text
                    if "_xlfn.ANCHORARRAY" in text:
                        text = convert_formula_text(text, sheetname, index)
                        changed_anchor += 1
                    new_text = strip_xlfn_prefixes(text)
                    if new_text != v.text:
                        c.value = ArrayFormula(ref=v.ref, text=new_text)
                        if new_text != text:
                            changed_prefix += 1
                elif isinstance(v, str) and v.startswith("="):
                    text = v
                    if "_xlfn.ANCHORARRAY" in text:
                        text = convert_formula_text(text, sheetname, index)
                        changed_anchor += 1
                    new_text = strip_xlfn_prefixes(text)
                    if new_text != v:
                        c.value = new_text
                        if new_text != text:
                            changed_prefix += 1

    print(f"Celdas con ANCHORARRAY resuelto: {changed_anchor}")
    print(f"Celdas con prefijo _xlfn./_xlws. quitado: {changed_prefix}")

    changed_map = apply_map_replacements(wb)
    print(f"Formulas MAP reemplazadas por BYROW/CHOOSE: {changed_map}")

    # verificacion: no debe quedar ningun ANCHORARRAY, prefijo _xlfn. ni MAP(
    remaining = 0
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                txt = v.text if isinstance(v, ArrayFormula) else (v if isinstance(v, str) else None)
                if txt and ("_xlfn." in txt or "ANCHORARRAY" in txt or re.search(r"(?<![A-Z])MAP\(", txt)):
                    remaining += 1
                    print("QUEDA SIN CONVERTIR:", sheetname, c.coordinate, txt[:120])
    if remaining:
        print(f"ADVERTENCIA: quedaron {remaining} ocurrencias sin convertir. No se guarda.")
        sys.exit(1)

    wb.save(PATH)
    print(f"Guardado OK: {PATH}")


if __name__ == "__main__":
    main()

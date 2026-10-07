"""
Crea la pestana MEMO_EQUIPO (antes de README) con un resumen de una pagina
para compartir con el equipo: que hay nuevo, 3 advertencias de
interpretacion, y el aviso de que Excel y la copia de Google Sheets estan
desincronizadas por diseno.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
SHEET_NAME = "MEMO_EQUIPO"

BLUE = "FFD9E1F2"
YELLOW = "FFFFF2CC"


def set_row(ws, row, text, *, bold=False, size=11, italic=False, fill=None, wrap=True, height=None):
    cell = ws.cell(row=row, column=2, value=text)
    cell.font = Font(bold=bold, size=size, italic=italic)
    if fill:
        for col in range(2, 7):
            ws.cell(row=row, column=col).fill = PatternFill(fill_type="solid", fgColor=fill)
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    cell.alignment = Alignment(wrap_text=wrap, vertical="top")
    if height:
        ws.row_dimensions[row].height = height
    return cell


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    if SHEET_NAME in wb.sheetnames:
        del wb[SHEET_NAME]
    ws = wb.create_sheet(SHEET_NAME, 0)  # 0 = primera pestaña, antes de README

    ws.column_dimensions["A"].width = 2
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20
    ws.sheet_view.showGridLines = False

    r = 2
    set_row(ws, r, "MEMO — Base consolidada por RUC (cooperativas café/cacao)", bold=True, size=15)
    ws.row_dimensions[r].height = 22
    r += 1
    set_row(ws, r, "Actualizado 24/09/2026. Resumen de una página para el equipo — el detalle técnico completo está en la pestaña README.", italic=True, size=10)
    r += 2

    set_row(ws, r, "QUÉ HAY NUEVO HOY", bold=True, size=12, fill=BLUE)
    r += 1
    for line in [
        "• RUC_FOB_DUA — número de DUA y canal de control (verde/rojo/naranja) por operación de exportación.",
        "• RESUMEN_EXPORT_VENTAS_PRODUCTO — exportación sobre ventas, agrupado por producto (café, cacao, etc.).",
        "• DIM_ORGANIZACION!AW:AX — n_duas y n_duas_canal_no_verde, nuevas columnas por RUC.",
        "• 3 scripts nuevos de reconexión con fuentes externas (ver README, sección de scripts) — permiten refrescar RUC_FOB, RUC_OECE/RNP/RUC_PPA y RUC_FOB_DUA sin repetir el trabajo de hoy.",
    ]:
        set_row(ws, r, line, size=10.5)
        r += 1
    r += 1

    set_row(ws, r, "3 ADVERTENCIAS DE INTERPRETACIÓN — leer antes de sacar conclusiones", bold=True, size=12, fill=YELLOW)
    r += 1
    warn = [
        ("1. ventas_totales NO es el tamaño real del negocio.",
         "Es autoreportado al programa (b2), no la facturación total de la cooperativa. Para exportadores "
         "verificados, el FOB de SUNAT llegó a ser 5-30x mayor que las ventas_totales reportadas del mismo "
         "año. No usar ventas_totales como proxy de tamaño para cooperativas con exportación fuerte."),
        ("2. Canal ROJO/NARANJA en aduanas es control de riesgo rutinario, no una señal de irregularidad.",
         "Solo 6 de 1,704 DUA salieron ROJO (0.35%) — muestra demasiado chica para ver un patrón real. No "
         "nombrar cooperativas específicas por esto en ningún reporte o comunicación."),
        ("3. La partida arancelaria (qué producto exacto se exportó) sigue sin poder obtenerse.",
         "El detalle línea por línea de cada DUA está protegido por captcha (servlet SgCDUI2 de SUNAT). Lo "
         "que sí tenemos es el nivel RUC-año-mes-aduana-agente-país, no el producto/partida por embarque."),
    ]
    for title, body in warn:
        set_row(ws, r, title, bold=True, size=10.5)
        r += 1
        set_row(ws, r, body, size=10.5)
        r += 1
    r += 1

    set_row(ws, r, "EXCEL vs. GOOGLE SHEETS — son dos archivos separados", bold=True, size=12, fill=BLUE)
    r += 1
    set_row(ws, r,
        "BASE CONSOLIDADA.xlsx (Excel, con fórmulas en vivo) es la fuente de verdad. "
        "BASE CONSOLIDADA_GSheets.xlsx es una conversión de un solo sentido para que abra en Google Sheets "
        "sin errores (Sheets no soporta algunas funciones de array dinámico de Excel). "
        "Si alguien edita la copia de Sheets, ese cambio NO vuelve al Excel — son archivos independientes. "
        "Antes de editar algo que debe persistir, confirmar que se está trabajando sobre el Excel original.",
        size=10.5, height=60)
    r += 2

    set_row(ws, r, "Dudas o hallazgos nuevos → agregar como nota en la pestaña correspondiente y avisar al resto del equipo antes de publicar cualquier cifra hacia afuera.", italic=True, size=10)

    wb.save(XLSX_PATH)
    print(f"Guardado OK: pestaña {SHEET_NAME} creada como primera pestaña (antes de README).")
    print("orden de pestañas:", wb.sheetnames[:3])


if __name__ == "__main__":
    main()

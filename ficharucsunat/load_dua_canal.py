"""
Carga dua_canal_final.csv (numero de DUA + canal de control por operacion de
RUC_FOB) como pestana nueva RUC_FOB_DUA en BASE CONSOLIDADA.xlsx.

No es la partida arancelaria (esa sigue bloqueada por captcha, ver hallazgo
del 23/09/2026) -- es el numero de DUA y el canal (verde/rojo/naranja) de
cada embarque, obtenido sin captcha via el endpoint "buscarListadoDuas" de
Aduanet. Ver dua_canal_connector.py para el detalle del endpoint.
"""
import csv

import openpyxl

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
CSV_PATH = "dua_canal_final.csv"
SHEET_NAME = "RUC_FOB_DUA"


def main():
    with open(CSV_PATH, encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    print(f"filas a cargar: {len(rows)}")

    wb = openpyxl.load_workbook(XLSX_PATH, data_only=False)
    if SHEET_NAME in wb.sheetnames:
        raise SystemExit(f"{SHEET_NAME} ya existe, no se sobreescribe automaticamente")

    ws = wb.create_sheet(SHEET_NAME)
    headers = ["ruc", "anio", "mes", "aduana_codigo", "agente_codigo", "pais_codigo",
               "dua_numero", "fecha_embarque", "canal", "url_detalle", "fecha_consulta"]
    ws.append(headers)

    for r in rows:
        ws.append([r[h] for h in headers])

    ws.freeze_panes = "A2"
    from openpyxl.styles import Font, PatternFill
    header_fill = PatternFill(fill_type="solid", fgColor="FFD9E1F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    wb.save(XLSX_PATH)
    print(f"Guardado OK: pestana {SHEET_NAME} con {len(rows)} filas")


if __name__ == "__main__":
    main()

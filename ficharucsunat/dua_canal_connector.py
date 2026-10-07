"""
Ingesta de numero de DUA + canal de control (verde/rojo/naranja) por cada
operacion ya cargada en RUC_FOB, consultando "buscarListadoDuas" de
Aduanet-SUNAT. Es el mismo dominio que exportaciones-fob-connector.ts
(appsperu/apps/identidad-fiscal), confirmado en vivo sin bloqueo ni captcha
para ESTE nivel (el detalle linea-por-linea de la DUA en si, con partida
arancelaria, SI pide captcha via SgCDUI2 — no se intenta, ver hallazgo del
23/09/2026).

Endpoint: GET http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias
  ?accion=buscarListadoDuas&CG_consulta=1&CG_regimen=40&CG_ruc=4{ruc}
  &CG_aduana={aduana_codigo}&CG_agente={agente_codigo}&CG_mes={mes}
  &CG_anio={anio}&CG_pais={pais_codigo}

CG_ruc lleva un "4" antepuesto al RUC de 11 digitos (tipo de documento),
confirmado en vivo revisando el onclick de jsDetalleDUA en la pagina
resumen. CG_regimen=40 fijo (exportacion definitiva, mismo valor que ya usa
exportaciones-fob-connector.ts).

Una combinacion ruc-anio-mes-aduana-agente-pais puede tener MAS de una DUA
(embarques distintos con los mismos codigos) — se guardan todas, no se
asume 1 a 1 con la fila de RUC_FOB.

Uso:
  python dua_canal_connector.py [--limit N] [--out CSV]
Lee "BASE CONSOLIDADA.xlsx" (pestana RUC_FOB) y escribe un CSV con el
resultado (no toca el xlsx — eso se hace despues, revisando el CSV).
"""
import argparse
import csv
import re
import sys
import time
from datetime import datetime, timezone
from urllib.request import urlopen, Request
from urllib.parse import urlencode

import openpyxl
from bs4 import BeautifulSoup

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
BASE_URL = "http://www.aduanet.gob.pe/cl-ad-itconsultadwh/ieITS01Alias"
DELAY_S = 0.4


def read_ruc_fob_rows():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["RUC_FOB"]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        vals = [c.value for c in row]
        if vals[0] is None:
            continue
        rows.append(dict(zip(headers, vals)))
    return rows


def build_url(ruc, anio, mes, aduana_codigo, agente_codigo, pais_codigo):
    params = {
        "accion": "buscarListadoDuas",
        "CG_consulta": "1",
        "CG_regimen": "40",
        "CG_ruc": f"4{ruc}",
        "CG_aduana": str(aduana_codigo),
        "CG_agente": str(agente_codigo),
        "CG_mes": str(mes),
        "CG_anio": str(anio),
        "CG_pais": pais_codigo,
    }
    return f"{BASE_URL}?{urlencode(params)}"


def fetch(url, timeout=20):
    req = Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urlopen(req, timeout=timeout) as resp:
        raw = resp.read()
    return raw.decode("latin1")


def parse_duas(html):
    soup = BeautifulSoup(html, "html.parser")
    out = []
    for tr in soup.select("tr.bg"):
        tds = tr.find_all("td")
        if len(tds) < 5:
            continue
        aduana = tds[0].get_text(strip=True)
        anio = tds[1].get_text(strip=True)
        a = tds[2].find("a")
        if a is None:
            continue
        dua_numero = a.get_text(strip=True)
        href = a.get("href", "")
        fecha_embarque = tds[3].get_text(strip=True)
        canal = tds[4].get_text(strip=True).upper()
        if not aduana.isdigit() or not anio.isdigit():
            continue
        out.append(
            {
                "aduana_codigo": aduana,
                "anio_dua": anio,
                "dua_numero": dua_numero,
                "fecha_embarque": fecha_embarque,
                "canal": canal,
                "url_detalle": href,
            }
        )
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=None, help="Procesar solo las primeras N filas (pruebas)")
    ap.add_argument("--out", default="dua_canal.csv")
    ap.add_argument("--resume-from", type=int, default=0, help="Saltar las primeras N filas de RUC_FOB (para reanudar)")
    args = ap.parse_args()

    rows = read_ruc_fob_rows()
    if args.limit:
        rows = rows[args.resume_from : args.resume_from + args.limit]
    else:
        rows = rows[args.resume_from :]

    fieldnames = [
        "ruc",
        "anio",
        "mes",
        "aduana_codigo",
        "agente_codigo",
        "pais_codigo",
        "dua_numero",
        "fecha_embarque",
        "canal",
        "url_detalle",
        "fecha_consulta",
    ]

    mode = "a" if args.resume_from else "w"
    write_header = not (args.resume_from and mode == "a")
    n_ok, n_sin_dua, n_err = 0, 0, 0
    t0 = time.time()

    with open(args.out, mode, newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        if write_header:
            w.writeheader()

        for i, r in enumerate(rows, start=1):
            url = build_url(r["ruc"], r["anio"], r["mes"], r["aduana_codigo"], r["agente_codigo"], r["pais_codigo"])
            fecha_consulta = datetime.now(timezone.utc).isoformat()
            try:
                html = fetch(url)
                duas = parse_duas(html)
                if not duas:
                    n_sin_dua += 1
                for d in duas:
                    w.writerow(
                        {
                            "ruc": r["ruc"],
                            "anio": r["anio"],
                            "mes": r["mes"],
                            "aduana_codigo": r["aduana_codigo"],
                            "agente_codigo": r["agente_codigo"],
                            "pais_codigo": r["pais_codigo"],
                            "dua_numero": d["dua_numero"],
                            "fecha_embarque": d["fecha_embarque"],
                            "canal": d["canal"],
                            "url_detalle": d["url_detalle"],
                            "fecha_consulta": fecha_consulta,
                        }
                    )
                n_ok += 1
            except Exception as e:
                n_err += 1
                print(f"[{i}/{len(rows)}] ERROR ruc={r['ruc']} anio={r['anio']} mes={r['mes']}: {e}", flush=True)

            if i % 50 == 0 or i == len(rows):
                elapsed = time.time() - t0
                print(
                    f"[{i}/{len(rows)}] ok={n_ok} sin_dua={n_sin_dua} err={n_err} "
                    f"({elapsed:.0f}s, {elapsed/i:.2f}s/fila)",
                    flush=True,
                )

            time.sleep(DELAY_S)

    print(f"Listo. ok={n_ok} sin_dua={n_sin_dua} err={n_err}. CSV: {args.out}")


if __name__ == "__main__":
    sys.exit(main())

"""
Reintento puntual de dua_canal_connector.py para las filas de RUC_FOB del
historico 2024 que no tenian agente_codigo (venia vacio en el archivo
aportado por el usuario, solo tenian agente_nombre). Se resuelve el codigo
por nombre normalizado contra las filas que SI lo tienen (del resto de
RUC_FOB), y se reintenta solo esas combinaciones.
"""
import csv
import re
import time
from datetime import datetime, timezone

import openpyxl
from dua_canal_connector import build_url, fetch, parse_duas, DELAY_S

XLSX_PATH = "BASE CONSOLIDADA.xlsx"
OUT_CSV = "dua_canal_2024_reintento.csv"


def norm(s):
    return re.sub(r"\s+", " ", (s or "").strip().upper())


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["RUC_FOB"]
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        vals = [c.value for c in row]
        if vals[0] is None:
            continue
        rows.append(dict(zip(headers, vals)))

    lookup = {}
    for r in rows:
        if r["agente_codigo"]:
            lookup.setdefault(norm(r["agente_nombre"]), set()).add(r["agente_codigo"])

    faltantes = [r for r in rows if not r["agente_codigo"]]
    print(f"filas sin agente_codigo: {len(faltantes)}")

    resueltas = []
    no_resueltas = []
    for r in faltantes:
        codes = lookup.get(norm(r["agente_nombre"]))
        if codes and len(codes) == 1:
            r2 = dict(r)
            r2["agente_codigo"] = next(iter(codes))
            resueltas.append(r2)
        else:
            no_resueltas.append(r)

    print(f"resueltas por nombre: {len(resueltas)}  sin resolver: {len(no_resueltas)}")
    for r in no_resueltas:
        print("  SIN RESOLVER:", r["agente_nombre"])

    fieldnames = [
        "ruc", "anio", "mes", "aduana_codigo", "agente_codigo", "pais_codigo",
        "dua_numero", "fecha_embarque", "canal", "url_detalle", "fecha_consulta",
    ]
    n_ok, n_sin_dua, n_err = 0, 0, 0
    with open(OUT_CSV, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for i, r in enumerate(resueltas, start=1):
            url = build_url(r["ruc"], r["anio"], r["mes"], r["aduana_codigo"], r["agente_codigo"], r["pais_codigo"])
            fecha_consulta = datetime.now(timezone.utc).isoformat()
            try:
                html = fetch(url)
                duas = parse_duas(html)
                if not duas:
                    n_sin_dua += 1
                for d in duas:
                    w.writerow({
                        "ruc": r["ruc"], "anio": r["anio"], "mes": r["mes"],
                        "aduana_codigo": r["aduana_codigo"], "agente_codigo": r["agente_codigo"],
                        "pais_codigo": r["pais_codigo"], "dua_numero": d["dua_numero"],
                        "fecha_embarque": d["fecha_embarque"], "canal": d["canal"],
                        "url_detalle": d["url_detalle"], "fecha_consulta": fecha_consulta,
                    })
                n_ok += 1
            except Exception as e:
                n_err += 1
                print(f"[{i}/{len(resueltas)}] ERROR: {e}", flush=True)
            if i % 50 == 0 or i == len(resueltas):
                print(f"[{i}/{len(resueltas)}] ok={n_ok} sin_dua={n_sin_dua} err={n_err}", flush=True)
            time.sleep(DELAY_S)

    print(f"Listo. ok={n_ok} sin_dua={n_sin_dua} err={n_err}. CSV: {OUT_CSV}")
    print(f"Nombres de agente que quedaron sin poder resolverse (no se reintentaron): {len(no_resueltas)}")


if __name__ == "__main__":
    main()

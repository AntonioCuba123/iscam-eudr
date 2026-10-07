# -*- coding: utf-8 -*-
import openpyxl, csv

# Load SUNAT
wb = openpyxl.load_workbook(
    r'C:\Users\acuba\.minimax\v2\assets\2026\09\17\22-12-33-164-asset_20260917-221233-164_99e4013bb7eb_b08b6316-SUNAT exportaciones_2010_2025 (1).xlsx',
    read_only=True, data_only=True)
ws = wb.active

sunat_rucs = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    ruc = str(r[6]).strip() if r[6] else None
    if not ruc: continue
    anio = int(r[7]) if r[7] else None
    fob = float(r[5]) if r[5] else 0
    pais = r[4] if r[4] else ''
    if ruc not in sunat_rucs:
        sunat_rucs[ruc] = {'fob': 0, 'ops': 0, 'paises': set(), 'years': set()}
    sunat_rucs[ruc]['fob'] += fob
    sunat_rucs[ruc]['ops'] += 1
    sunat_rucs[ruc]['paises'].add(pais)
    if anio: sunat_rucs[ruc]['years'].add(anio)
wb.close()

print(f'Exportadores en SUNAT Aduanas: {len(sunat_rucs)}')

# Load universo
universo_rucs = set()
try:
    with open(r'C:\Users\acuba\eudrperu\data\universo_rucs.csv', 'r', encoding='utf-8') as f:
        reader = csv.DictReader(f)
        for row in reader:
            universo_rucs.add(row['ruc'].strip())
except:
    # Try matrix xlsx
    wb2 = openpyxl.load_workbook(r'C:\Users\acuba\eudrperu\MATRIZ_IP_IV_v5_EUDRPERU.xlsx', read_only=True, data_only=True)
    ws2 = wb2['MATRIZ']
    for row in ws2.iter_rows(min_row=2, values_only=True):
        if row[0]:  # RUC column A
            universo_rucs.add(str(row[0]).strip())
    wb2.close()

print(f'Universo RUCs: {len(universo_rucs)}')

match = universo_rucs & set(sunat_rucs.keys())
only_sunat = set(sunat_rucs.keys()) - universo_rucs
print(f'Match universo-SUNAT: {len(match)}')
print(f'En SUNAT fuera universo: {len(only_sunat)}')

# Top 10
top = sorted(sunat_rucs.items(), key=lambda x: x[1]['fob'], reverse=True)[:10]
print('\nTop 10 por FOB total:')
for ruc, d in top:
    yr = sorted(d['years'])
    label = f'{min(yr)}-{max(yr)}' if yr else '?'
    print(f'  {ruc}: {d["fob"]:>15,.0f} USD | {d["ops"]:>4} ops | {len(d["paises"]):>2} paises | {label}')

# New exporters not in universo
if only_sunat:
    print(f'\n{len(only_sunat)} exportadores en SUNAT fuera del universo actual:')
    for ruc in sorted(only_sunat):
        d = sunat_rucs[ruc]
        yr = sorted(d['years'])
        label = f'{min(yr)}-{max(yr)}' if yr else '?'
        print(f'  {ruc}: {d["fob"]:>12,.0f} USD | {d["ops"]} ops | {label}')

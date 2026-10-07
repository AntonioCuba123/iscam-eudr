import openpyxl
import re

wb = openpyxl.load_workbook('C:/Users/acuba/eudrperu/MATRIZ_IP_IV_v17_EUDRPERU.xlsx')
ws = wb['MATRIZ']

# Las formulas BD/BE/BF en fila 7 se borraron en el swap
# Escribir formulas para TODAS las filas de datos (desde fila 7 hasta max_row)

max_row = ws.max_row
print('Datos desde fila 7 hasta ' + str(max_row) + ' = ' + str(max_row - 6) + ' orgs')

# Formulas
formulas = {
    'BD': '=0',
    'BE': '=IFERROR((IF(ISNUMBER(T{r}),MIN(T{r}/30,1),0)+IF(ISNUMBER(P{r}),MIN(P{r}/500,1),0))/2,0)',
    'BF': '=IFERROR(IF(OR(U{r}<>"",V{r}="SÍ"),IF(ISNUMBER(L{r}),(1*MIN(L{r}/10000,1))^0.5,0),IF(ISNUMBER(L{r}),(0*MIN(L{r}/10000,1))^0.5,0)),0)',
}

formula_fill = openpyxl.styles.PatternFill('solid', fgColor='E2EFDA')
formula_font = openpyxl.styles.Font(name='Calibri', size=10)
formula_align = openpyxl.styles.Alignment(horizontal='center', vertical='center')

count = 0
for r in range(7, max_row + 1):
    for col, template in formulas.items():
        cell = ws[col + str(r)]
        formula = template.replace('{r}', str(r))
        cell.value = formula
        cell.font = formula_font
        cell.fill = formula_fill
        cell.alignment = formula_align
        cell.number_format = '0.000'
        count += 1

print('Formulas escritas: ' + str(count))
print('BE7:', ws['BE7'].value)
print('BF7:', ws['BF7'].value)
print('BE8:', ws['BE8'].value)
print('BF8:', ws['BF8'].value)
print('Fila 7 RUC:', ws['B7'].value, '| Nombre:', ws['C7'].value)

wb.save('C:/Users/acuba/eudrperu/MATRIZ_IP_IV_v17_EUDRPERU.xlsx')
print('Guardado v17')

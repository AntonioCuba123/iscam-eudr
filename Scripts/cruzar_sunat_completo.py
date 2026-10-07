# -*- coding: utf-8 -*-
"""Cruce completo SUNAT Aduanas vs universo EUDRPERU + enriquecimiento"""
import openpyxl, csv, requests, time

SUNAT_FILE = r'C:\Users\acuba\.minimax\v2\assets\2026\09\17\22-12-33-164-asset_20260917-221233-164_99e4013bb7eb_b08b6316-SUNAT exportaciones_2010_2025 (1).xlsx'
MATRIX_FILE = r'C:\Users\acuba\eudrperu\MATRIZ_IP_IV_v5_EUDRPERU.xlsx'
API_URL = 'http://localhost:4006/api/ruc/'

# ── 1. Cargar SUNAT ──────────────────────────────────────────────────────────
print('Cargando SUNAT Aduanas...')
wb = openpyxl.load_workbook(SUNAT_FILE, read_only=True, data_only=True)
ws = wb.active

sunat_data = {}   # ruc -> {fob_total, ops, paises, years, fob_anio: {}, name}
for r in ws.iter_rows(min_row=2, values_only=True):
    ruc_raw = r[6]
    if not ruc_raw: continue
    ruc = str(ruc_raw).strip()
    anio = int(r[7]) if r[7] else None
    fob = float(r[5]) if r[5] else 0
    pais = str(r[4]).strip() if r[4] else ''
    name_raw = r[0] if r[0] else ''
    name = str(name_raw).split('-', 1)[0].strip() if name_raw else ''

    if ruc not in sunat_data:
        sunat_data[ruc] = {
            'name': name, 'fob_total': 0, 'ops': 0,
            'paises': set(), 'years': set(), 'fob_anio': {}
        }
    sunat_data[ruc]['fob_total'] += fob
    sunat_data[ruc]['ops'] += 1
    sunat_data[ruc]['paises'].add(pais)
    if anio: sunat_data[ruc]['years'].add(anio)
    if anio:
        sunat_data[ruc]['fob_anio'][anio] = sunat_data[ruc]['fob_anio'].get(anio, 0) + fob

wb.close()
print(f'  {len(sunat_data)} exportadores cargados')

# ── 2. Cargar universo actual ────────────────────────────────────────────────
print('Cargando universo actual...')
universo = {}
try:
    wb2 = openpyxl.load_workbook(MATRIX_FILE, data_only=True)
    ws2 = wb2['MATRIZ']
    headers = [c.value for c in ws2[1]]
    for row in ws2.iter_rows(min_row=2, values_only=True):
        ruc = str(row[0]).strip() if row[0] else ''
        if ruc:
            universo[ruc] = dict(zip(headers, row))
    wb2.close()
    print(f'  {len(universo)} organizaciones en universo')
except Exception as e:
    print(f'  Error cargando matriz: {e}')
    universo = {}

# ── 3. Identificar RUCs nuevos vs existentes ────────────────────────────────
existentes = set(universo.keys())
todos_sunat = set(sunat_data.keys())
nuevos = todos_sunat - existentes
matched = todos_sunat & existentes

print(f'  Existentes en universo: {len(matched)}')
print(f'  Nuevos (no en universo): {len(nuevos)}')

# ── 4. Validar y filtrar nuevos (café/cacao) ────────────────────────────────
print('\nValidando nuevos contra API SUNAT...')
keywords = ['cafe', 'café', 'cacao', 'cocoa', 'cafetera', 'cafetalera',
            'cacaotera', 'cacaotera', 'coffee', 'cocoa']
nuevos_validos = {}

for ruc in sorted(nuevos):
    d = sunat_data[ruc]
    name = d['name'].lower()
    # Filtrar por keyword en nombre
    is_cafe_cacao = any(k.lower() in name for k in keywords)
    if not is_cafe_cacao:
        print(f'  SKIP {ruc}: no cafe/cacao | name={d["name"]!r}')
        continue

    # Validar RUC en API
    try:
        resp = requests.get(API_URL + ruc, timeout=5)
        if resp.status_code == 200:
            data_api = resp.json()
            estado = data_api.get('estado', '')
            razon = data_api.get('razon_social', '')
        else:
            estado = 'NO_VALIDADO'
            razon = ''
    except Exception as e:
        estado = 'ERROR_API:' + str(e)[:30]
        razon = ''

    nuevos_validos[ruc] = {
        'name_sunat': d['name'],
        'razon_sunat': razon,
        'estado': estado,
        'fob_total': d['fob_total'],
        'ops': d['ops'],
        'paises': len(d['paises']),
        'years': sorted(d['years']),
        'fob_anio': d['fob_anio'],
        'is_cafe_cacao': is_cafe_cacao
    }
    print(f'  OK   {ruc}: {razon or d["name"]} | FOB={d["fob_total"]:,.0f} | {estado}')

print(f'\nNuevos válidos café/cacao: {len(nuevos_validos)}')

# ── 5. Enriquecer universo existente con series SUNAT ────────────────────────
print('\nEnriqueciendo universo existente con datos SUNAT...')
enriquecidos = {}
for ruc in matched:
    d = sunat_data[ruc]
    yr = sorted(d['years'])
    # Serie FOB por año
    serie_fob = '; '.join([f'{a}:{d["fob_anio"].get(a,0):.0f}' for a in yr])
    enriquecidos[ruc] = {
        'FOB_TOTAL_SUNAT': d['fob_total'],
        'SUNAT_OPS': d['ops'],
        'SUNAT_PAISES': len(d['paises']),
        'SUNAT_PAIS_LISTA': '; '.join(sorted(d['paises'])),
        'SUNAT_ANIO_MIN': min(yr) if yr else '',
        'SUNAT_ANIO_MAX': max(yr) if yr else '',
        'SUNAT_SERIE_FOB': serie_fob,
        'SUNAT_2010': d['fob_anio'].get(2010, 0),
        'SUNAT_2011': d['fob_anio'].get(2011, 0),
        'SUNAT_2012': d['fob_anio'].get(2012, 0),
        'SUNAT_2013': d['fob_anio'].get(2013, 0),
        'SUNAT_2014': d['fob_anio'].get(2014, 0),
        'SUNAT_2015': d['fob_anio'].get(2015, 0),
        'SUNAT_2016': d['fob_anio'].get(2016, 0),
        'SUNAT_2017': d['fob_anio'].get(2017, 0),
        'SUNAT_2018': d['fob_anio'].get(2018, 0),
        'SUNAT_2019': d['fob_anio'].get(2019, 0),
        'SUNAT_2020': d['fob_anio'].get(2020, 0),
        'SUNAT_2021': d['fob_anio'].get(2021, 0),
        'SUNAT_2022': d['fob_anio'].get(2022, 0),
        'SUNAT_2023': d['fob_anio'].get(2023, 0),
        'SUNAT_2024': d['fob_anio'].get(2024, 0),
        'SUNAT_2025': d['fob_anio'].get(2025, 0),
    }

print(f'  {len(enriquecidos)} organizaciones enriquecidas')

# ── 6. Guardar resultados ────────────────────────────────────────────────────
# Guardar nuevos candidatos
nuevos_path = r'C:\Users\acuba\eudrperu\data\sunat_nuevos_candidatos.csv'
with open(nuevos_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=[
        'ruc', 'name_sunat', 'razon_sunat', 'estado_sunat', 'fob_total',
        'ops', 'paises', 'anio_min', 'anio_max', 'is_cafe_cacao'
    ])
    writer.writeheader()
    for ruc, d in sorted(nuevos_validos.items()):
        yr = d['years']
        writer.writerow({
            'ruc': ruc, 'name_sunat': d['name_sunat'],
            'razon_sunat': d['razon_sunat'], 'estado_sunat': d['estado'],
            'fob_total': d['fob_total'], 'ops': d['ops'],
            'paises': d['paises'],
            'anio_min': min(yr) if yr else '',
            'anio_max': max(yr) if yr else '',
            'is_cafe_cacao': d['is_cafe_cacao']
        })
print(f'\nNuevos candidatos guardados: {nuevos_path}')

# Guardar enriquecimiento
enr_path = r'C:\Users\acuba\eudrperu\data\sunat_enriquecimiento.csv'
with open(enr_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.DictWriter(f, fieldnames=list(next(iter(enriquecidos.values())).keys()))
    writer.writeheader()
    for ruc, d in sorted(enriquecidos.items()):
        writer.writerow({'RUC': ruc, **d})
print(f'Enriquecimiento guardado: {enr_path}')

print('\n--- RESUMEN ---')
print(f'Match universo-SUNAT: {len(matched)}')
print(f'Nuevos cafe/cacao encontrados: {len(nuevos_validos)}')
total_nuevos_fob = sum(d['fob_total'] for d in nuevos_validos.values())
print(f'FOB nuevos (USD): {total_nuevos_fob:,.0f}')

import openpyxl
import re

wb = openpyxl.load_workbook('C:/Users/acuba/eudrperu/MATRIZ_IP_IV_v14_EUDRPERU.xlsx')
ws = wb['MATRIZ']

# Tras insert_rows(1):
# - Datos原来的fila6 ahora en fila7
# - Las formulas BD/BE/BF原来的fila6 ahora en fila7 pero sus refs internas siguen en {r}=6
# Hay que cambiarlas a {r}=7, {r}=8, ... {r}=264+1=265

max_row = ws.max_row
print(f'Corrigiendo formulas BD/BE/BF para filas 7 a {max_row}')

def fix_formula(formula, delta=1):
    """Incrementa todas las referencias de fila en el formula en delta."""
    def replace_row(m):
        return m.group(1) + str(int(m.group(2)) + delta) + m.group(3)
    return re.sub(r'(\b[A-Z]+)(\d+)(\b)', replace_row, formula)

count_bd = 0
count_be = 0
count_bf = 0

for r in range(7, max_row + 1):
    for col in ['BD', 'BE', 'BF']:
        cell = ws[col + str(r)]
        if cell.value and isinstance(cell.value, str) and cell.value.startswith('='):
            old = cell.value
            new = fix_formula(cell.value, delta=1)
            if old != new:
                cell.value = new
                if col == 'BD': count_bd += 1
                elif col == 'BE': count_be += 1
                elif col == 'BF': count_bf += 1

print(f'Formulas corregidas: BD={count_bd}, BE={count_be}, BF={count_bf}')

# Verificar
for r in [7, 8, 9]:
    print(f'  Fila {r} - BE: {ws["BE"+str(r)].value}')
    print(f'  Fila {r} - BF: {ws["BF"+str(r)].value}')

# Verificar que headers en fila 5 esten bien
print(f'  Fila 5 headers (B,C,D): {ws["B5"].value} | {ws["C5"].value} | {ws["D5"].value}')
print(f'  Fila 7 datos: RUC={ws["B7"].value} | Nombre={ws["C7"].value}')

wb.save('C:/Users/acuba/eudrperu/MATRIZ_IP_IV_v15_EUDRPERU.xlsx')
print('Guardado v15 con formulas corregidas')

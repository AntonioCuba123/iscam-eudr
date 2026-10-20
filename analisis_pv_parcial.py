# -*- coding: utf-8 -*-
import openpyxl, collections, statistics

wb = openpyxl.load_workbook(r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx', data_only=True)
m = wb['MATRIZ_174']
vals, byaq, extrem = [], collections.defaultdict(list), []
for r in range(3, 177):
    s, aq, ap = m.cell(r, 19).value, m.cell(r, 43).value, m.cell(r, 42).value
    if isinstance(s, (int, float)):
        vals.append(s); byaq[aq].append(s)
        extrem.append((s, aq, r, m.cell(r, 2).value))
print('PVE score (S) — n =', len(vals))
print('  min %.1f | max %.1f | media %.1f | mediana %.1f' % (min(vals), max(vals),
      statistics.mean(vals), statistics.median(vals)))
print('  por cobertura (AQ):', {k: f'{len(v)} orgs, media {statistics.mean(v):.1f}'
      for k, v in sorted(byaq.items())})
extrem.sort(reverse=True)
print('  top 5:', [(f'{s:.1f}', f'AQ={aq}') for s, aq, r, n in extrem[:5]])
print('  bottom 5:', [(f'{s:.1f}', f'AQ={aq}') for s, aq, r, n in extrem[-5:]])
print('  (S = 100 en', sum(1 for s in vals if s == 100), 'organizaciones; S = 0 en',
      sum(1 for s in vals if s == 0), ')')
# de qué componente vive cada valor extremo
r0 = extrem[0][2]
print('  ejemplo top (fila', r0, '): M =', m.cell(r0, 13).value, '| O =', m.cell(r0, 15).value,
      '| W:Z =', [m.cell(r0, c).value for c in (23, 24, 25, 26)])

# -*- coding: utf-8 -*-
import openpyxl, warnings, re
warnings.filterwarnings('ignore')
wb = openpyxl.load_workbook(r'C:\Users\acuba\eudrperu\PLANTILLA_INDICE_PVE_ISCAM.xlsx')
rx = re.compile(r'MATRIZ_174!\$?O\d|SCORECARDS!\$?[A-Z]+')
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and re.search(r'MATRIZ_174!\$?O\d', v):
                print(ws.title, c.coordinate, '|', v[:230])
print('--- MATRIZ_174 fila 2 cabeceras completas ---')
m = wb['MATRIZ_174']
from openpyxl.utils import get_column_letter as gl
for c in range(1, m.max_column + 3):
    v = m.cell(2, c).value
    if v is not None:
        print(gl(c), '=', repr(str(v)[:60]))

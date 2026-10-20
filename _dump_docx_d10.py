# -*- coding: utf-8 -*-
import docx
d = docx.Document(r'C:\Users\acuba\eudrperu\ISCAM\Metodologia - Instrumento 2 e Indice ISCAM.docx')
print('parrafos', len(d.paragraphs), 'tablas', len(d.tables))
for i, p in enumerate(d.paragraphs):
    t = p.text
    if 'D10' in t or 'IP5' in t or 'asesor' in t.lower() or 'Certificaci' in t:
        print(f'[{i:03d}]', repr(t[:400]))
print('=== tablas ===')
for ti, tb in enumerate(d.tables):
    for ri, row in enumerate(tb.rows):
        line = ' | '.join(c.text.replace('\n', ' ')[:60] for c in row.cells)
        if 'IP5' in line or 'D10' in line or 'Cert' in line:
            print(f'T{ti} r{ri}:', line[:300])
